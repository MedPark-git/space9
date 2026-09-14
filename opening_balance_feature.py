import hashlib
import io
import json
import shutil
import uuid
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from flask import abort, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook as XlsxWorkbook, load_workbook

EXPECTED_HEADERS = ["구분", "금융기관", "계좌번호", "계좌명", "통화", "기준일", "기초잔액", "비고"]
BASE_DATE = date(2026, 1, 1)
SOURCE_DATE = date(2025, 12, 31)
DEFAULT_RECON_DATE = date(2026, 8, 31)
_registered = False


def _decimal(value):
    if value is None or value == "":
        raise ValueError("기초잔액이 비어 있습니다. 0원 계좌도 0을 입력해 주세요.")
    if isinstance(value, bool):
        raise ValueError("기초잔액 숫자 형식이 아닙니다.")
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"기초잔액 숫자 형식 오류: {value}") from exc


def _base_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        from datetime import timedelta
        try:
            return date(1899, 12, 30) + timedelta(days=int(value))
        except Exception:
            return None
    s = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _hash_headers(headers):
    return hashlib.sha256("|".join(headers).encode("utf-8")).hexdigest()


def _parse_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    if "기초잔액입력" not in wb.sheetnames:
        raise ValueError("[기초잔액입력] 시트를 찾을 수 없습니다.")
    ws = wb["기초잔액입력"]
    normalized = [str(ws.cell(5, c).value or "").strip() for c in range(1, 9)]
    if normalized != EXPECTED_HEADERS:
        raise ValueError("기초잔액 Excel 양식의 Header가 기준 양식과 일치하지 않습니다.")
    rows, errors, seen, totals = [], [], set(), {}
    for r in range(6, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, 9)]
        if not any(v not in (None, "") for v in values):
            continue
        kind = str(values[0] or "").strip(); bank = str(values[1] or "").strip(); account = str(values[2] or "").strip(); account_name = str(values[3] or "").strip(); currency = str(values[4] or "").strip().upper(); base_date = _base_date(values[5]); note = str(values[7] or "").strip()
        row_errors = []
        if kind not in {"예금", "MMT"}: row_errors.append("구분은 예금 또는 MMT만 가능합니다.")
        if not bank: row_errors.append("금융기관이 비어 있습니다.")
        if not account: row_errors.append("계좌번호가 비어 있습니다.")
        if not currency: row_errors.append("통화가 비어 있습니다.")
        if base_date != BASE_DATE: row_errors.append("기준일은 2026-01-01이어야 합니다.")
        try:
            opening = _decimal(values[6])
        except ValueError as exc:
            opening = None; row_errors.append(str(exc))
        key = (kind, bank, account, currency)
        if key in seen: row_errors.append("동일 구분·금융기관·계좌·통화가 중복되어 있습니다.")
        seen.add(key)
        if row_errors:
            errors.append({"row": r, "message": " / ".join(row_errors)}); continue
        totals[currency] = totals.get(currency, Decimal("0")) + opening
        rows.append({"source_row_number": r, "kind": kind, "bank": bank, "account_number": account, "account_name": account_name, "currency": currency, "base_date": BASE_DATE.isoformat(), "opening_balance": str(opening), "note": note})
    wb.close()
    if not rows and not errors:
        raise ValueError("등록할 기초잔액 행을 찾지 못했습니다.")
    return {"rows": rows, "errors": errors, "totals": {k: str(v) for k, v in sorted(totals.items())}, "structure_hash": _hash_headers(normalized)}


def _select_bank_reported(txs, currency):
    if not txs: return None
    latest_date = max(t.transaction_date for t in txs); same = [t for t in txs if t.transaction_date == latest_date]
    if currency != "KRW":
        return min(same, key=lambda t: t.source_row_number)
    return max(same, key=lambda t: (t.transaction_time or time.min, t.source_row_number))


def _select_mmt_reported(txs):
    if not txs: return None
    latest_date = max(t.transaction_date for t in txs); same = [t for t in txs if t.transaction_date == latest_date]
    return max(same, key=lambda t: (t.transaction_time or time.min, t.source_row_number))


def register():
    global _registered
    if _registered: return
    import app as core
    app, db = core.app, core.db

    class MmtOpeningBalance(db.Model):
        __tablename__ = "mmt_opening_balances"
        id = db.Column(db.Uuid(as_uuid=True), primary_key=True, default=uuid.uuid4)
        base_date = db.Column(db.Date, nullable=False, index=True)
        financial_institution = db.Column(db.String(80), nullable=False)
        account_number = db.Column(db.String(100), nullable=False)
        account_name = db.Column(db.String(200))
        currency = db.Column(db.String(10), nullable=False, default="KRW")
        opening_balance = db.Column(db.Numeric(28, 6), nullable=False)
        status = db.Column(db.String(30), nullable=False, default="confirmed_manual")
        calculation_method = db.Column(db.String(500), nullable=False)
        source_batch_id = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("import_batches.id"))
        created_by = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("users.id"), nullable=False)
        created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow)
        updated_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow, onupdate=core.utcnow)
        source_batch = db.relationship("ImportBatch")
        __table_args__ = (db.UniqueConstraint("base_date", "financial_institution", "account_number", "currency", name="uq_mmt_opening_balance_identity"),)
    core.MmtOpeningBalance = MmtOpeningBalance

    def opening_preview_path(token): return core.PENDING_DIR / f"opening_{token}.json"
    def opening_xlsx_path(token): return core.PENDING_DIR / f"opening_{token}.xlsx"
    def load_opening_preview(token):
        path = opening_preview_path(token)
        if not path.exists(): abort(404)
        with path.open("r", encoding="utf-8") as f: data = json.load(f)
        if data.get("created_by") != str(current_user.id) and current_user.role != "admin": abort(403)
        return data

    @login_required
    def opening_balances_view():
        rows = core.OpeningBalance.query.filter_by(base_date=BASE_DATE).join(core.BankAccount).order_by(core.BankAccount.financial_institution, core.BankAccount.account_number, core.BankAccount.currency).all()
        mmt_rows = MmtOpeningBalance.query.filter_by(base_date=BASE_DATE).order_by(MmtOpeningBalance.financial_institution, MmtOpeningBalance.account_number, MmtOpeningBalance.currency).all()
        manual_count = sum(1 for r in rows if r.status == "confirmed_manual") + len(mmt_rows); auto_count = sum(1 for r in rows if r.status == "confirmed"); review_count = sum(1 for r in rows if r.status == "needs_review")
        return render_template("opening_balances.html", rows=rows, mmt_rows=mmt_rows, base_date=BASE_DATE, manual_count=manual_count, auto_count=auto_count, review_count=review_count)
    app.view_functions["opening_balances"] = opening_balances_view

    @app.get("/data/opening-balances/template")
    @login_required
    def opening_balance_template():
        out = XlsxWorkbook(); ws = out.active; ws.title = "기초잔액입력"
        ws.append(["MedPark 2026년 기초잔액 일괄등록 양식"]); ws.append(["G열 [기초잔액]을 입력하세요. 0원 계좌도 0을 입력해야 합니다."]); ws.append([]); ws.append([]); ws.append(EXPECTED_HEADERS)
        accounts = core.BankAccount.query.order_by(core.BankAccount.financial_institution, core.BankAccount.account_number, core.BankAccount.currency).all()
        for a in accounts:
            ob = core.OpeningBalance.query.filter_by(base_date=BASE_DATE, bank_account_id=a.id).first(); ws.append(["예금", a.financial_institution, a.account_number, a.account_name or "", a.currency, BASE_DATE, float(ob.opening_balance) if ob and ob.opening_balance is not None else None, ""])
        mmt_keys = db.session.query(core.MmtTransaction.financial_institution, core.MmtTransaction.account_number, core.MmtTransaction.account_name, core.MmtTransaction.currency).distinct().order_by(core.MmtTransaction.financial_institution, core.MmtTransaction.account_number).all()
        for bank, account, account_name, currency in mmt_keys:
            ob = MmtOpeningBalance.query.filter_by(base_date=BASE_DATE, financial_institution=bank, account_number=account, currency=currency).first(); ws.append(["MMT", bank, account, account_name or "", currency, BASE_DATE, float(ob.opening_balance) if ob else None, ""])
        for cell in ws[5]: cell.font = cell.font.copy(bold=True)
        ws.freeze_panes = "A6"
        for col, width in {"A":10,"B":12,"C":25,"D":32,"E":10,"F":14,"G":20,"H":30}.items(): ws.column_dimensions[col].width = width
        bio = io.BytesIO(); out.save(bio); bio.seek(0)
        return send_file(bio, as_attachment=True, download_name="MedPark_2026_기초잔액_일괄등록_양식.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.post("/data/opening-balances/upload")
    @core.roles_required("admin", "editor")
    def opening_balance_upload():
        uploaded = request.files.get("file")
        if not uploaded or not uploaded.filename: flash("기초잔액 Excel 파일을 선택해 주세요.", "error"); return redirect(url_for("opening_balances"))
        if Path(uploaded.filename).suffix.lower() != ".xlsx": flash("기초잔액은 지정된 .xlsx 양식만 업로드할 수 있습니다.", "error"); return redirect(url_for("opening_balances"))
        token = uuid.uuid4().hex; path = opening_xlsx_path(token); uploaded.save(path)
        try:
            parsed = _parse_xlsx(path); counts = {"total": len(parsed["rows"]), "bank": 0, "mmt": 0, "new": 0, "new_account": 0, "update": 0, "error": len(parsed["errors"])}
            for row in parsed["rows"]:
                if row["kind"] == "예금":
                    counts["bank"] += 1; account = core.BankAccount.query.filter_by(financial_institution=row["bank"], account_number=row["account_number"], currency=row["currency"]).first()
                    if not account: row["action"] = "new_account"; counts["new_account"] += 1
                    else:
                        ob = core.OpeningBalance.query.filter_by(base_date=BASE_DATE, bank_account_id=account.id).first(); row["action"] = "update" if ob else "new"; counts[row["action"]] += 1
                else:
                    counts["mmt"] += 1; ob = MmtOpeningBalance.query.filter_by(base_date=BASE_DATE, financial_institution=row["bank"], account_number=row["account_number"], currency=row["currency"]).first(); row["action"] = "update" if ob else "new"; counts[row["action"]] += 1
            data = {"token": token, "created_by": str(current_user.id), "original_filename": uploaded.filename, "file_sha256": core.file_sha256(path), "structure_hash": parsed["structure_hash"], "rows": parsed["rows"], "errors": parsed["errors"], "totals": parsed["totals"], "counts": counts}
            with opening_preview_path(token).open("w", encoding="utf-8") as f: json.dump(data, f, ensure_ascii=False)
            core.audit("opening_balance_preview_created", user=current_user, target_type="opening_balance_import", target_id=token, detail=f"rows={counts['total']} errors={counts['error']}"); core.commit_or_rollback(); return redirect(url_for("opening_balance_preview", token=token))
        except Exception as exc:
            path.unlink(missing_ok=True); core.app.logger.exception("Opening balance preview failed"); flash(f"기초잔액 파일 검증 실패: {exc}", "error"); return redirect(url_for("opening_balances"))

    @app.get("/data/opening-balances/preview/<token>")
    @login_required
    def opening_balance_preview(token): return render_template("opening_balance_preview.html", data=load_opening_preview(token))

    @app.post("/data/opening-balances/confirm/<token>")
    @core.roles_required("admin", "editor")
    def opening_balance_confirm(token):
        data = load_opening_preview(token); xlsx_path = opening_xlsx_path(token)
        if not xlsx_path.exists() or core.file_sha256(xlsx_path) != data["file_sha256"]: flash("원본 파일 검증에 실패했습니다. 다시 업로드해 주세요.", "error"); return redirect(url_for("opening_balances"))
        parsed = _parse_xlsx(xlsx_path)
        if parsed["errors"]: flash("오류가 있는 파일은 등록할 수 없습니다.", "error"); return redirect(url_for("opening_balance_preview", token=token))
        stored_file = None
        try:
            batch_code = core.next_batch_code(); stored_file = core.IMPORT_DIR / f"{batch_code}_opening_balance.xlsx"; shutil.copy2(xlsx_path, stored_file)
            batch = core.ImportBatch(batch_code=batch_code, file_type="opening_balance", original_filename=data["original_filename"], stored_path=str(stored_file), file_sha256=data["file_sha256"], structure_hash=parsed["structure_hash"], query_start_date=BASE_DATE, query_end_date=BASE_DATE, total_rows=len(parsed["rows"]), new_rows=len(parsed["rows"]), duplicate_rows=0, review_rows=0, error_rows=0, status="confirmed", created_by=current_user.id, confirmed_at=core.utcnow()); db.session.add(batch); db.session.flush(); new_count = update_count = 0
            for row in parsed["rows"]:
                opening = Decimal(row["opening_balance"]); action = "new"
                if row["kind"] == "예금":
                    account = core.BankAccount.query.filter_by(financial_institution=row["bank"], account_number=row["account_number"], currency=row["currency"]).first()
                    if not account:
                        account = core.BankAccount(financial_institution=row["bank"], account_number=row["account_number"], account_name=row["account_name"], currency=row["currency"], account_type="deposit"); db.session.add(account); db.session.flush(); action = "new_account"
                    elif row["account_name"] and account.account_name != row["account_name"]: account.account_name = row["account_name"]
                    ob = core.OpeningBalance.query.filter_by(base_date=BASE_DATE, bank_account_id=account.id).first()
                    if ob: action = "update"
                    else: ob = core.OpeningBalance(base_date=BASE_DATE, bank_account_id=account.id, currency=row["currency"], created_by=current_user.id); db.session.add(ob)
                    ob.currency = row["currency"]; ob.source_date = SOURCE_DATE; ob.source_balance = None; ob.opening_balance = opening; ob.status = "confirmed_manual"; ob.calculation_method = "Excel 일괄등록 - 사용자 확정(2025.12.31 기말 / 2026.01.01 기초)"; ob.source_batch_id = batch.id; ob.source_transaction_id = None
                else:
                    ob = MmtOpeningBalance.query.filter_by(base_date=BASE_DATE, financial_institution=row["bank"], account_number=row["account_number"], currency=row["currency"]).first()
                    if ob: action = "update"
                    else: ob = MmtOpeningBalance(base_date=BASE_DATE, financial_institution=row["bank"], account_number=row["account_number"], currency=row["currency"], created_by=current_user.id); db.session.add(ob)
                    ob.account_name = row["account_name"]; ob.opening_balance = opening; ob.status = "confirmed_manual"; ob.calculation_method = "Excel 일괄등록 - 사용자 확정(2025.12.31 기말 / 2026.01.01 기초)"; ob.source_batch_id = batch.id
                db.session.add(core.ImportRawRow(import_batch_id=batch.id, source_row_number=row["source_row_number"], row_status=action, message=row.get("note") or None, raw_data=row)); update_count += 1 if action == "update" else 0; new_count += 0 if action == "update" else 1
            batch.new_rows = new_count; core.audit("opening_balances_imported", user=current_user, target_type="import_batch", target_id=batch.id, detail=f"{batch_code} rows={len(parsed['rows'])} new={new_count} update={update_count}"); db.session.commit(); xlsx_path.unlink(missing_ok=True); opening_preview_path(token).unlink(missing_ok=True)
            flash(f"기초잔액 등록 완료: {len(parsed['rows'])}계좌/통화 · 신규 {new_count} · 갱신 {update_count}", "success"); return redirect(url_for("opening_balances"))
        except Exception:
            db.session.rollback()
            if stored_file: Path(stored_file).unlink(missing_ok=True)
            core.app.logger.exception("Opening balance confirm failed"); flash("기초잔액 등록 중 오류가 발생해 DB 반영을 롤백했습니다.", "error"); return redirect(url_for("opening_balance_preview", token=token))

    @core.roles_required("admin", "editor")
    def generate_opening_balances_safe():
        accounts = core.BankAccount.query.order_by(core.BankAccount.financial_institution, core.BankAccount.account_number, core.BankAccount.currency).all(); confirmed = review = skipped = 0
        for account in accounts:
            ob = core.OpeningBalance.query.filter_by(base_date=BASE_DATE, bank_account_id=account.id).first()
            if ob and ob.status == "confirmed_manual": skipped += 1; continue
            txs = core.BankTransaction.query.filter_by(account_id=account.id, transaction_date=SOURCE_DATE).order_by(core.BankTransaction.source_row_number).all(); helper = [{"id":str(t.id), "transaction_time":t.transaction_time.isoformat() if t.transaction_time else None, "balance":str(t.balance), "deposit":str(t.deposit_amount), "withdrawal":str(t.withdrawal_amount)} for t in txs]; final_data, _status, method = core.select_final_balance(helper)
            if not ob: ob = core.OpeningBalance(base_date=BASE_DATE, bank_account_id=account.id, currency=account.currency, created_by=current_user.id); db.session.add(ob)
            ob.source_date = SOURCE_DATE; ob.calculation_method = method
            if final_data:
                tx = next(t for t in txs if str(t.id) == final_data["id"]); ob.source_balance = tx.balance; ob.opening_balance = tx.balance; ob.status = "confirmed"; ob.source_batch_id = tx.import_batch_id; ob.source_transaction_id = tx.id; confirmed += 1
            else:
                ob.source_balance = None; ob.opening_balance = None; ob.status = "needs_review"; ob.source_batch_id = None; ob.source_transaction_id = None; review += 1
        core.audit("opening_balances_generated", user=current_user, target_type="opening_balance", target_id=BASE_DATE.isoformat(), detail=f"confirmed={confirmed}, review={review}, manual_skipped={skipped}")
        if core.commit_or_rollback(): flash(f"자동산출 완료: 자동 확인 {confirmed} / 확인필요 {review} / Excel 확정 유지 {skipped}", "success")
        else: flash("기초잔액 자동산출에 실패했습니다.", "error")
        return redirect(url_for("opening_balances"))
    app.view_functions["generate_opening_balances"] = generate_opening_balances_safe

    @app.get("/data/opening-balances/reconcile")
    @login_required
    def opening_balance_reconcile():
        cutoff_text = (request.args.get("cutoff") or DEFAULT_RECON_DATE.isoformat()).strip()
        try: cutoff = datetime.strptime(cutoff_text, "%Y-%m-%d").date()
        except ValueError: cutoff = DEFAULT_RECON_DATE
        results, totals = [], {}
        bank_obs = core.OpeningBalance.query.filter_by(base_date=BASE_DATE).join(core.BankAccount).order_by(core.BankAccount.financial_institution, core.BankAccount.account_number, core.BankAccount.currency).all()
        for ob in bank_obs:
            txs = core.BankTransaction.query.filter(core.BankTransaction.account_id == ob.bank_account_id, core.BankTransaction.transaction_date >= BASE_DATE, core.BankTransaction.transaction_date <= cutoff).all(); deposits = sum((Decimal(t.deposit_amount) for t in txs), Decimal("0")); withdrawals = sum((Decimal(t.withdrawal_amount) for t in txs), Decimal("0")); opening = Decimal(ob.opening_balance or 0); calculated = opening + deposits - withdrawals; reported_tx = _select_bank_reported(txs, ob.currency); reported = Decimal(reported_tx.balance) if reported_tx else None; diff = calculated - reported if reported is not None else None
            results.append({"kind":"예금", "bank":ob.account.financial_institution, "account":ob.account.account_number, "account_name":ob.account.account_name, "currency":ob.currency, "opening":opening, "deposits":deposits, "withdrawals":withdrawals, "calculated":calculated, "reported":reported, "difference":diff, "reported_date":reported_tx.transaction_date if reported_tx else None, "match": diff == 0 if diff is not None else False})
        for ob in MmtOpeningBalance.query.filter_by(base_date=BASE_DATE).order_by(MmtOpeningBalance.financial_institution, MmtOpeningBalance.account_number).all():
            txs = core.MmtTransaction.query.filter(core.MmtTransaction.financial_institution == ob.financial_institution, core.MmtTransaction.account_number == ob.account_number, core.MmtTransaction.currency == ob.currency, core.MmtTransaction.transaction_date >= BASE_DATE, core.MmtTransaction.transaction_date <= cutoff).all(); deposits = sum((Decimal(t.deposit_amount) for t in txs), Decimal("0")); withdrawals = sum((Decimal(t.withdrawal_amount) for t in txs), Decimal("0")); opening = Decimal(ob.opening_balance or 0); calculated = opening + deposits - withdrawals; reported_tx = _select_mmt_reported(txs); reported = Decimal(reported_tx.balance) if reported_tx else None; diff = calculated - reported if reported is not None else None
            results.append({"kind":"MMT", "bank":ob.financial_institution, "account":ob.account_number, "account_name":ob.account_name, "currency":ob.currency, "opening":opening, "deposits":deposits, "withdrawals":withdrawals, "calculated":calculated, "reported":reported, "difference":diff, "reported_date":reported_tx.transaction_date if reported_tx else None, "match": diff == 0 if diff is not None else False})
        for r in results:
            t = totals.setdefault(r["currency"], {"opening":Decimal("0"), "deposits":Decimal("0"), "withdrawals":Decimal("0"), "calculated":Decimal("0"), "reported":Decimal("0"), "count":0, "mismatch":0})
            for field in ("opening","deposits","withdrawals","calculated"): t[field] += r[field]
            if r["reported"] is not None: t["reported"] += r["reported"]
            t["count"] += 1; t["mismatch"] += 0 if r["match"] else 1
        return render_template("opening_balance_reconciliation.html", results=results, totals=totals, cutoff=cutoff, mismatch_count=sum(1 for r in results if not r["match"]))

    _registered = True
