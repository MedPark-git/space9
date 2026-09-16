import io, uuid
from datetime import date
from decimal import Decimal, InvalidOperation
from flask import Response, abort, flash, redirect, request, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Font

import app as pkg
import phase2_feature as p2
import phase2_enhancements as p2e
import phase21_fx_core as fx
from fx_provider_smb import SOURCE, RATE_TYPE

core=pkg.core;db=pkg.db;app=pkg.app;REGISTERED=False

def audit(action,detail):core.audit(action,user=current_user,target_type='fx_rate',detail=detail)

def _panel(day,snapshot=None):
    info=(snapshot or {}).get('fx_rates_used') if snapshot else fx.rate_info(day)
    if not info:return "<section class='panel'><h3>환율 상태</h3><p class='muted'>사용 가능한 환율이 없습니다.</p></section>"
    rs=''.join(f"<tr><td>{c}</td><td class='num'>{v['rate']}</td><td class='num'>{v['unit_amount']} {c}</td><td>{v['source']}</td><td>{v['source_date']}</td></tr>" for c,v in sorted(info.items()))
    return f"<section class='panel'><h3>적용 환율</h3><table><thead><tr><th>통화</th><th>매매기준율</th><th>적용단위</th><th>Source</th><th>고시일</th></tr></thead><tbody>{rs}</tbody></table></section>"

def _inject(resp,day,snapshot=None):
    r=app.make_response(resp)
    if 'text/html' in r.headers.get('Content-Type',''):
        h=r.get_data(as_text=True);r.set_data(h.replace('</main>',_panel(day,snapshot)+'</main>'))
    return r

def register():
    global REGISTERED
    if REGISTERED:return
    with app.app_context():fx.migrate();fx.patch_snapshot()

    @login_required
    def manage():
        day=p2._d(request.args.get('date'),p2._latest_transaction_date() or date.today());rs=fx.rows(day);required=fx.required_currencies()
        body="""<div class='p2head'><div><h2>환율관리</h2><p class='muted'>서울외국환중개 매매기준율 자동조회가 기본입니다. 수기등록은 예외/Fallback으로 유지합니다.</p></div></div>
        <section class='panel'><h3>서울외국환중개 환율 조회</h3><form method='post' action='{{ url_for("phase21_fx_lookup") }}' class='filters'><input type='hidden' name='csrf_token' value='{{ csrf_token() }}'><label>기준일<input type='date' name='base_date' value='{{ day }}' required></label><label>조회범위<select name='scope'><option value='required'>MedPark 사용통화 우선</option><option value='all'>전체 주요통화</option></select></label><button>서울외국환중개 환율 조회</button></form><p class='muted'>현재 외화계좌 통화: {{ required|join(', ') if required else '없음' }} · 화면 진입만으로 외부 사이트를 호출하지 않습니다.</p></section>
        <section class='panel'><h3>등록 환율 — {{ day }}</h3><div class='scroll'><table><thead><tr><th>통화</th><th class='num'>환율</th><th class='num'>단위</th><th>종류</th><th>Source</th><th>고시일</th><th>Active</th><th>조회/등록일</th></tr></thead><tbody>{% for r in rates %}<tr><td>{{ r.currency }}</td><td class='num'>{{ r.rate }}</td><td class='num'>{{ r.unit_amount if r.unit_amount is not none else '-' }}</td><td>{{ r.rate_type or '기존자료' }}</td><td>{{ r.source }}</td><td>{{ r.source_date or '-' }}</td><td>{{ 'Y' if r.is_active else 'N' }}</td><td>{{ (r.retrieved_at or r.created_at)|kst }}</td></tr>{% else %}<tr><td colspan='8' class='empty'>등록된 환율이 없습니다.</td></tr>{% endfor %}</tbody></table></div></section>
        <details class='panel'><summary><b>수기등록 / 예외등록</b></summary><form method='post' action='{{ url_for("phase2_fx_rate_add") }}' class='filters' style='margin-top:12px'><input type='hidden' name='csrf_token' value='{{ csrf_token() }}'><label>기준일<input type='date' name='base_date' value='{{ day }}' required></label><label>통화<input name='currency' placeholder='USD' required></label><label>매매기준율<input type='number' step='0.00000001' name='rate' required></label><label>적용단위<input type='number' step='0.00000001' name='unit_amount' value='1' required></label><label>Source<input name='source' value='수기등록' required></label><button>수기 환율 등록</button></form></details>"""
        return p2._render('환율관리',body,day=day,rates=rs,required=required)

    @core.roles_required('admin','editor')
    def lookup():
        day=p2._d(request.form.get('base_date'));scope=request.form.get('scope') or 'required';previous=request.form.get('previous')=='1'
        if not day:abort(400)
        currs=fx.required_currencies() if scope=='required' else sorted(set(fx.required_currencies()+fx.ALL))
        if not currs:flash('조회할 외화 통화가 없습니다.','error');return redirect(url_for('phase2_fx_rates',date=day.isoformat()))
        try:
            raw=fx.PROVIDER.lookup(day,currs,previous);token=uuid.uuid4().hex;out=[];errs=[]
            for r in raw:
                if r.get('status')!='lookup_ok':errs.append(r);out.append(r);continue
                st,e=fx.compare(day,r);r['register_status']=st;r['existing']={'rate':str(e['rate']),'unit_amount':str(e.get('unit_amount')) if e.get('unit_amount') is not None else None,'source':e.get('source'),'source_date':str(e.get('source_date') or e.get('base_date'))} if e else None
                prev=fx.previous_effective(day,r['currency']);eff=Decimal(r['rate'])/Decimal(r['unit_amount']);r['warning']=('전일/직전 등록 환율 대비 변동폭이 비정상적으로 큽니다.' if prev and prev>0 and (eff/prev>=5 or eff/prev<=Decimal('0.2')) else None);out.append(r)
            data={'token':token,'created_by':str(current_user.id),'base_date':day.isoformat(),'scope':scope,'previous':previous,'rows':out,'errors':errs,'retrieved_at':core.utcnow().isoformat()};fx.save_preview(token,data);audit('fx_rate_auto_lookup',f"date={day} source={SOURCE} currencies={','.join(currs)} previous={previous}");db.session.commit();return redirect(url_for('phase21_fx_preview',token=token))
        except Exception as e:
            db.session.rollback();audit('fx_rate_auto_lookup_failed',f'date={day} source={SOURCE} error={str(e)[:250]}');db.session.commit();flash(f'서울외국환중개 환율 조회에 실패했습니다. {str(e)[:180]}','error');return redirect(url_for('phase2_fx_rates',date=day.isoformat()))

    @login_required
    def preview(token):
        d=fx.load_preview(token,current_user.id,current_user.role=='admin')
        if not d:abort(404)
        body="""<div class='p2head'><div><h2>서울외국환중개 환율 조회결과</h2><p class='muted'>기준일 {{ d.base_date }} · 조회만으로 DB 저장되지 않습니다.</p></div><a class='btnlink secondary' href='{{ url_for("phase2_fx_rates") }}?date={{ d.base_date }}'>취소</a></div>
        <section class='panel scroll'><table><thead><tr><th>통화</th><th class='num'>환율</th><th class='num'>적용단위</th><th>고시일</th><th>상태</th><th>기존값</th><th>경고</th><th>변경승인</th></tr></thead><tbody>{% for r in d.rows %}<tr><td>{{ r.currency }}</td><td class='num'>{{ r.rate if r.rate is defined else '-' }}</td><td class='num'>{% if r.unit_amount is defined %}{{ r.unit_amount }} {{ r.currency }}{% else %}-{% endif %}</td><td>{{ r.source_date if r.source_date is defined else '-' }}</td><td>{% if r.register_status is defined %}{{ {'new':'신규','registered':'기등록','change_review':'변경 확인 필요'}.get(r.register_status,r.register_status) }}{% else %}{{ r.status }}{% endif %}</td><td>{% if r.existing %}{{ r.existing.rate }} / {{ r.existing.source }}{% else %}-{% endif %}</td><td>{{ r.warning or r.message or '-' }}</td><td>{% if r.register_status=='change_review' %}<input form='cf' type='checkbox' name='update_currency' value='{{ r.currency }}'> 조회값으로 갱신{% else %}-{% endif %}</td></tr>{% endfor %}</tbody></table></section>
        {% if d.errors and not d.previous %}<section class='panel'><div class='warnbox'>선택일자에 일부/전체 고시환율이 없습니다. 직전 영업일을 자동 적용하지 않습니다.</div><form method='post' action='{{ url_for("phase21_fx_lookup") }}'><input type='hidden' name='csrf_token' value='{{ csrf_token() }}'><input type='hidden' name='base_date' value='{{ d.base_date }}'><input type='hidden' name='scope' value='{{ d.scope }}'><input type='hidden' name='previous' value='1'><button>직전 고시일 환율 조회</button></form></section>{% endif %}
        <section class='panel'><h3>조회 결과 등록</h3><p class='muted'>기존값과 차이가 있는 통화는 위에서 갱신 승인을 선택해야 변경됩니다.</p><form id='cf' method='post' action='{{ url_for("phase21_fx_confirm",token=d.token) }}'><input type='hidden' name='csrf_token' value='{{ csrf_token() }}'><button>조회 결과 등록</button></form></section>"""
        return p2._render('환율 조회결과',body,d=d)

    @core.roles_required('admin','editor')
    def confirm(token):
        d=fx.load_preview(token,current_user.id,current_user.role=='admin')
        if not d:abort(404)
        day=date.fromisoformat(d['base_date']);allowed=set(request.form.getlist('update_currency'));ins=upd=skip=0;changes=[]
        try:
            for r in d['rows']:
                if r.get('status')!='lookup_ok':continue
                st=r.get('register_status')
                if st=='registered':skip+=1;continue
                if st=='change_review' and r['currency'] not in allowed:skip+=1;continue
                existing=next((x for x in fx.rows(day,r['currency']) if x.get('source')==SOURCE and (x.get('rate_type') or RATE_TYPE)==RATE_TYPE),None)
                if existing:old=existing['rate'];fx.update_auto(existing['id'],day,r);upd+=1;changes.append(f"{r['currency']}:{old}->{r['rate']}")
                else:fx.insert_auto(day,r,current_user.id);ins+=1
            audit('fx_rate_auto_saved',f'date={day} source={SOURCE} inserted={ins} updated={upd} skipped={skip}')
            if changes:audit('fx_rate_auto_updated',f"date={day} changes={'|'.join(changes)}")
            if d.get('previous'):audit('fx_rate_previous_source_date_used',f"base_date={day} source_dates={','.join(sorted({r.get('source_date','') for r in d['rows'] if r.get('status')=='lookup_ok'}))}")
            db.session.commit();snap=p2._snapshot(day,pkg.MmtOpeningBalance);audit('cash_journal_draft_recalculated',f"date={day} fx_total={snap.get('fx_krw_total')} ready={snap.get('ready')}");db.session.commit();fx.preview_path(token).unlink(missing_ok=True);flash(f'환율 등록 완료: 신규 {ins} / 갱신 {upd} / 유지·미등록 {skip}. 자금일보 Draft 재계산 완료.','success');return redirect(url_for('cash_journal',date=day.isoformat()))
        except Exception:
            db.session.rollback();core.app.logger.exception('FX confirm failed');flash('환율 등록 중 오류가 발생했습니다. DB에는 반영되지 않았습니다.','error');return redirect(url_for('phase21_fx_preview',token=token))

    @core.roles_required('admin','editor')
    def manual():
        day=p2._d(request.form.get('base_date'));c=(request.form.get('currency') or '').strip().upper();src=(request.form.get('source') or '수기등록').strip()
        try:r=Decimal(request.form.get('rate') or '0');u=Decimal(request.form.get('unit_amount') or '0')
        except InvalidOperation:r=u=Decimal('0')
        if not day or not c or r<=0 or u<=0:flash('기준일·통화·환율·적용단위를 확인해 주세요.','error');return redirect(url_for('phase2_fx_rates',date=(day or date.today()).isoformat()))
        rid,old=fx.save_manual(day,c,r,u,src,current_user.id);audit('fx_rate_manual_saved',f"date={day} currency={c} rate={r} unit={u} source={src} old={old.get('rate') if old else '-'}");db.session.commit();flash('수기 환율이 등록되었습니다.','success');return redirect(url_for('phase2_fx_rates',date=day.isoformat()))

    app.view_functions['phase2_fx_rates']=manage;app.view_functions['phase2_fx_rate_add']=manual
    app.add_url_rule('/reference/fx-rates/lookup','phase21_fx_lookup',lookup,methods=['POST']);app.add_url_rule('/reference/fx-rates/preview/<token>','phase21_fx_preview',preview,methods=['GET']);app.add_url_rule('/reference/fx-rates/confirm/<token>','phase21_fx_confirm',confirm,methods=['POST'])

    old_d=app.view_functions.get('dashboard');old_j=app.view_functions.get('cash_journal')
    if old_d:
        def dash(*a,**k):return _inject(old_d(*a,**k),p2._latest_transaction_date() or date.today())
        app.view_functions['dashboard']=dash
    if old_j:
        def journal(*a,**k):
            day=p2._d(request.args.get('date'),p2._latest_transaction_date() or date.today());ver=request.args.get('version');s=None
            if ver and str(ver).isdigit():
                v=p2.CashJournalVersion.query.filter_by(journal_date=day,version=int(ver)).first();s=v.snapshot_json if v else None
            return _inject(old_j(*a,**k),day,s)
        app.view_functions['cash_journal']=journal

    @login_required
    def excel():
        day=p2._d(request.args.get('date'),p2._latest_transaction_date() or date.today());ver=request.args.get('version');v=p2.CashJournalVersion.query.filter_by(journal_date=day,version=int(ver)).first() if ver and str(ver).isdigit() else None;s=v.snapshot_json if v else p2e._snapshot(day);wb=Workbook();ws=wb.active;ws.title='자금일보';ws.append(['MedPark 자금일보',day.isoformat()]);ws['A1'].font=Font(bold=True,size=16)
        for label,key in [('원화예금','krw_total'),('외화 원화환산','fx_krw_total'),('MMT','mmt_total'),('총 금융자금','total_financial'),('D+1 예상자금','expected')]:ws.append([label,Decimal(str(s.get(key,'0')))])
        ws.append([]);ws.append(['은행','계좌','통화','원통화잔액','환율','단위','원화환산','Source','고시일'])
        for x in s.get('accounts',[]):ws.append([x.get('bank'),x.get('account'),x.get('currency'),Decimal(x['balance']),Decimal(x['rate']) if x.get('rate') else None,Decimal(x['unit_amount']) if x.get('unit_amount') else None,Decimal(x['krw']) if x.get('krw') else None,x.get('rate_source'),x.get('rate_source_date')])
        sh=wb.create_sheet('환율정보');sh.append(['통화','매매기준율','적용단위','환율종류','Source','고시일'])
        for c,i in sorted((s.get('fx_rates_used') or {}).items()):sh.append([c,Decimal(i['rate']),Decimal(i['unit_amount']),i['rate_type'],i['source'],i['source_date']])
        out=io.BytesIO();wb.save(out);out.seek(0);return Response(out.getvalue(),mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename=MedPark_CashJournal_{day.isoformat()}.xlsx'})
    app.view_functions['phase2_cash_journal_excel']=excel;REGISTERED=True
