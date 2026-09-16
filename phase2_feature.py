import io
import json
import re
import sys
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

from flask import Response, abort, flash, redirect, render_template_string, request, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, or_, text

core = sys.modules.get("medpark_cash_core")
if core is None:
    raise RuntimeError("medpark_cash_core is not loaded")
db = core.db

BASE_DATE = date(2026, 1, 1)
PHASE2_LOCK = 874221940
REGISTERED = False

CATEGORY_NAMES = [
    "원부자재", "설비투자", "인허가", "품질비용", "연구개발", "과제지원", "기타비용(기술)",
    "상품구입", "국내전시회", "해외전시회", "광고선전비", "판매장려금", "기타비용(마케팅)",
    "인건비", "세금과공과", "이자비용", "원금상환", "임차료/관리비", "지급수수료", "여비교통비",
    "기타비용(경영)", "유형자산", "국내_덴탈 수금", "국내_메디컬 수금", "국내_에스테틱 수금",
    "해외_덴탈 수금", "해외_메디컬 수금", "기타수입", "과제수입", "운반비", "3공장", "특허", "보증금",
]


class BankAccountSetting(db.Model):
    __tablename__ = "bank_account_settings"
    id = db.Column(db.Uuid(as_uuid=True), primary_key=True, default=uuid.uuid4)
    bank_account_id = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("bank_accounts.id"), nullable=False, unique=True, index=True)
    include_in_cash_journal = db.Column(db.Boolean, nullable=False, default=True)
    display_order = db.Column(db.Integer, nullable=False, default=100)
    is_internal = db.Column(db.Boolean, nullable=False, default=True)
    note = db.Column(db.String(500))
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow)
    updated_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow, onupdate=core.utcnow)
    account = db.relationship("BankAccount")


class TransactionCategory(db.Model):
    __tablename__ = "transaction_categories"
    id = db.Column(db.Uuid(as_uuid=True), primary_key=True, default=uuid.uuid4)
    code = db.Column(db.String(20), nullable=False, unique=True, index=True)
    name = db.Column(db.String(100), nullable=False, unique=True, index=True)
    sort_order = db.Column(db.Integer, nullable=False, default=100)
    is_active = db.Column(db.Boolean, nullable=False, default=True)
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow)
    updated_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow, onupdate=core.utcnow)


class TransactionAnnotation(db.Model):
    __tablename__ = "transaction_annotations"
    id = db.Column(db.Uuid(as_uuid=True), primary_key=True, default=uuid.uuid4)
    transaction_id = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("bank_transactions.id"), nullable=False, unique=True, index=True)
    category_id = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("transaction_categories.id"))
    suggested_category_id = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("transaction_categories.id"))
    flow_type = db.Column(db.String(30), nullable=False, default="external")
    suggestion_reason = db.Column(db.String(500))
    classification_source = db.Column(db.String(30), nullable=False, default="manual")
    is_confirmed = db.Column(db.Boolean, nullable=False, default=False)
    confirmed_by = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("users.id"))
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow)
    updated_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow, onupdate=core.utcnow)
    transaction = db.relationship("BankTransaction")
    category = db.relationship("TransactionCategory", foreign_keys=[category_id])
    suggested_category = db.relationship("TransactionCategory", foreign_keys=[suggested_category_id])


class DailyBalanceCheck(db.Model):
    __tablename__ = "daily_balance_checks"
    id = db.Column(db.Uuid(as_uuid=True), primary_key=True, default=uuid.uuid4)
    bank_account_id = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("bank_accounts.id"), nullable=False, index=True)
    check_date = db.Column(db.Date, nullable=False, index=True)
    opening_balance = db.Column(db.Numeric(28, 6))
    cumulative_deposit = db.Column(db.Numeric(28, 6), nullable=False, default=0)
    cumulative_withdrawal = db.Column(db.Numeric(28, 6), nullable=False, default=0)
    calculated_balance = db.Column(db.Numeric(28, 6))
    bank_reported_balance = db.Column(db.Numeric(28, 6))
    difference = db.Column(db.Numeric(28, 6))
    status = db.Column(db.String(30), nullable=False, default="insufficient")
    checked_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow)
    account = db.relationship("BankAccount")
    __table_args__ = (
        db.UniqueConstraint("bank_account_id", "check_date", name="uq_daily_balance_account_date"),
    )


class CashPlan(db.Model):
    __tablename__ = "cash_plans"
    id = db.Column(db.Uuid(as_uuid=True), primary_key=True, default=uuid.uuid4)
    planned_date = db.Column(db.Date, nullable=False, index=True)
    direction = db.Column(db.String(20), nullable=False)
    bank_account_id = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("bank_accounts.id"))
    counterparty = db.Column(db.String(200))
    description = db.Column(db.String(500), nullable=False)
    category_id = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("transaction_categories.id"))
    currency = db.Column(db.String(10), nullable=False, default="KRW")
    amount = db.Column(db.Numeric(28, 6), nullable=False)
    status = db.Column(db.String(20), nullable=False, default="planned")
    note = db.Column(db.String(500))
    source = db.Column(db.String(30), nullable=False, default="manual")
    created_by = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("users.id"), nullable=False)
    updated_by = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("users.id"), nullable=False)
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow)
    updated_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow, onupdate=core.utcnow)
    account = db.relationship("BankAccount")
    category = db.relationship("TransactionCategory")


class FxRate(db.Model):
    __tablename__ = "fx_rates"
    id = db.Column(db.Uuid(as_uuid=True), primary_key=True, default=uuid.uuid4)
    base_date = db.Column(db.Date, nullable=False, index=True)
    currency = db.Column(db.String(10), nullable=False, index=True)
    rate = db.Column(db.Numeric(28, 8), nullable=False)
    round_no = db.Column(db.Integer, nullable=False, default=1)
    source = db.Column(db.String(100), nullable=False, default="하나은행")
    created_by = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("users.id"), nullable=False)
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow)
    __table_args__ = (
        db.UniqueConstraint("base_date", "currency", "round_no", "source", name="uq_fx_rate_daily"),
    )


class CashJournalVersion(db.Model):
    __tablename__ = "cash_journal_versions"
    id = db.Column(db.Uuid(as_uuid=True), primary_key=True, default=uuid.uuid4)
    journal_date = db.Column(db.Date, nullable=False, index=True)
    version = db.Column(db.Integer, nullable=False)
    status = db.Column(db.String(20), nullable=False)
    snapshot_json = db.Column(db.JSON, nullable=False)
    created_by = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("users.id"), nullable=False)
    reviewed_by = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("users.id"))
    confirmed_by = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("users.id"))
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow)
    confirmed_at = db.Column(db.DateTime(timezone=True))
    change_reason = db.Column(db.String(500))
    __table_args__ = (
        db.UniqueConstraint("journal_date", "version", name="uq_cash_journal_version"),
    )


def _d(value, default=None):
    if not value:
        return default
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    try:
        return datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except Exception:
        return default


def _dec(v):
    if v is None:
        return Decimal("0")
    return Decimal(str(v))


def _fmt(v, currency="KRW"):
    if v is None:
        return "-"
    d = _dec(v)
    if currency == "KRW":
        return f"{d:,.0f}"
    return f"{d:,.6f}".rstrip("0").rstrip(".")


def _norm_account(v):
    return re.sub(r"[^0-9A-Za-z]", "", str(v or "")).upper()


def _setting(account_id):
    return BankAccountSetting.query.filter_by(bank_account_id=account_id).first()


def _account_included(account):
    s = _setting(account.id)
    return account.is_active and (s.include_in_cash_journal if s else True)


def _account_internal(account):
    s = _setting(account.id)
    return s.is_internal if s else True


def _reported_tx(account, cutoff):
    q = core.BankTransaction.query.filter(core.BankTransaction.account_id == account.id, core.BankTransaction.transaction_date <= cutoff)
    if account.currency == "KRW":
        latest_date = db.session.query(func.max(core.BankTransaction.transaction_date)).filter(core.BankTransaction.account_id == account.id, core.BankTransaction.transaction_date <= cutoff).scalar()
        if not latest_date:
            return None
        same = q.filter(core.BankTransaction.transaction_date == latest_date).all()
        max_time = max((t.transaction_time for t in same if t.transaction_time is not None), default=None)
        if max_time is not None:
            same = [t for t in same if t.transaction_time == max_time]
        return min(same, key=lambda t: t.source_row_number) if same else None
    return q.order_by(core.BankTransaction.transaction_date.desc(), core.BankTransaction.source_row_number.asc()).first()


def _bank_balance(account, cutoff):
    ob = core.OpeningBalance.query.filter_by(base_date=BASE_DATE, bank_account_id=account.id).first()
    opening = _dec(ob.opening_balance) if ob and ob.opening_balance is not None else None
    dep, wd = db.session.query(func.coalesce(func.sum(core.BankTransaction.deposit_amount), 0), func.coalesce(func.sum(core.BankTransaction.withdrawal_amount), 0)).filter(core.BankTransaction.account_id == account.id, core.BankTransaction.transaction_date >= BASE_DATE, core.BankTransaction.transaction_date <= cutoff).one()
    calc = (opening + _dec(dep) - _dec(wd)) if opening is not None else None
    rt = _reported_tx(account, cutoff)
    reported = _dec(rt.balance) if rt else None
    if opening is None or reported is None:
        status = "insufficient"; diff = None
    else:
        diff = calc - reported; status = "ok" if diff == 0 else "review"
    return {"opening": opening, "deposit": _dec(dep), "withdrawal": _dec(wd), "calculated": calc, "reported": reported, "difference": diff, "status": status, "reported_date": rt.transaction_date if rt else None}


def _mmt_balance(mmt_opening_model, bank, account_no, currency, cutoff):
    ob = mmt_opening_model.query.filter_by(base_date=BASE_DATE, financial_institution=bank, account_number=account_no, currency=currency).first()
    opening = _dec(ob.opening_balance) if ob else None
    txs = core.MmtTransaction.query.filter(core.MmtTransaction.financial_institution == bank, core.MmtTransaction.account_number == account_no, core.MmtTransaction.currency == currency, core.MmtTransaction.transaction_date >= BASE_DATE, core.MmtTransaction.transaction_date <= cutoff).all()
    dep = sum((_dec(t.deposit_amount) for t in txs), Decimal("0")); wd = sum((_dec(t.withdrawal_amount) for t in txs), Decimal("0")); calc = opening + dep - wd if opening is not None else None
    if txs:
        latest_date = max(t.transaction_date for t in txs); same = [t for t in txs if t.transaction_date == latest_date]; max_time = max((t.transaction_time for t in same if t.transaction_time is not None), default=None)
        if max_time is not None: same = [t for t in same if t.transaction_time == max_time]
        last = min(same, key=lambda t: t.source_row_number); reported = _dec(last.balance)
    else:
        last = None; reported = None
    if opening is None or reported is None: status, diff = "insufficient", None
    else: diff = calc - reported; status = "ok" if diff == 0 else "review"
    return {"opening": opening, "deposit": dep, "withdrawal": wd, "calculated": calc, "reported": reported, "difference": diff, "status": status, "reported_date": last.transaction_date if last else None}


def _latest_transaction_date():
    return db.session.query(func.max(core.BankTransaction.transaction_date)).scalar()


def _latest_upload():
    return core.ImportBatch.query.filter(core.ImportBatch.file_type.in_(["krw", "fx"])).order_by(core.ImportBatch.confirmed_at.desc()).first()


def _all_mmt_keys(mmt_opening_model):
    keys = set()
    for r in db.session.query(core.MmtTransaction.financial_institution, core.MmtTransaction.account_number, core.MmtTransaction.currency).distinct().all(): keys.add(tuple(r))
    for r in mmt_opening_model.query.filter_by(base_date=BASE_DATE).all(): keys.add((r.financial_institution, r.account_number, r.currency))
    return sorted(keys)


def _ensure_category_seed():
    existing = {x.name: x for x in TransactionCategory.query.all()}
    for i, name in enumerate(CATEGORY_NAMES, start=1):
        code = f"C{i:03d}"; row = existing.get(name)
        if not row: db.session.add(TransactionCategory(code=code, name=name, sort_order=i, is_active=True))
        else: row.sort_order = i; row.is_active = True
    db.session.flush()


def _past_suggestion(tx):
    if not tx.description and not tx.counterparty: return None, None
    q = TransactionAnnotation.query.join(core.BankTransaction, TransactionAnnotation.transaction_id == core.BankTransaction.id).filter(TransactionAnnotation.is_confirmed.is_(True), TransactionAnnotation.category_id.is_not(None), core.BankTransaction.id != tx.id)
    if tx.counterparty: q = q.filter(core.BankTransaction.counterparty == tx.counterparty)
    if tx.description: q = q.filter(core.BankTransaction.description == tx.description)
    ann = q.order_by(TransactionAnnotation.updated_at.desc()).first()
    return (ann.category_id, "과거 동일 적요/상대방 확정패턴") if ann else (None, None)


def _ensure_annotations():
    existing = {x[0] for x in db.session.query(TransactionAnnotation.transaction_id).all()}
    internal_map = {_norm_account(a.account_number): a.id for a in core.BankAccount.query.all() if _account_internal(a)}
    q = core.BankTransaction.query
    if existing: q = q.filter(~core.BankTransaction.id.in_(existing))
    new_count = 0
    for tx in q.all():
        if tx.counter_account and _norm_account(tx.counter_account) in internal_map:
            flow = "internal_in" if _dec(tx.deposit_amount) > 0 else "internal_out" if _dec(tx.withdrawal_amount) > 0 else "internal"
        else:
            flow = "external_in" if _dec(tx.deposit_amount) > 0 else "external_out" if _dec(tx.withdrawal_amount) > 0 else "external"
        suggested_id, reason = _past_suggestion(tx)
        db.session.add(TransactionAnnotation(transaction_id=tx.id, flow_type=flow, suggested_category_id=suggested_id, suggestion_reason=reason, classification_source="rule" if suggested_id or flow.startswith("internal") else "manual")); new_count += 1
    return new_count


def _refresh_daily_checks(account_ids=None):
    q = core.BankAccount.query
    if account_ids: q = q.filter(core.BankAccount.id.in_(list(account_ids)))
    for account in q.all():
        ob = core.OpeningBalance.query.filter_by(base_date=BASE_DATE, bank_account_id=account.id).first(); opening = _dec(ob.opening_balance) if ob and ob.opening_balance is not None else None
        txs = core.BankTransaction.query.filter(core.BankTransaction.account_id == account.id, core.BankTransaction.transaction_date >= BASE_DATE).order_by(core.BankTransaction.transaction_date.asc(), core.BankTransaction.transaction_time.asc().nullsfirst(), core.BankTransaction.source_row_number.desc()).all()
        by_date = defaultdict(list)
        for tx in txs: by_date[tx.transaction_date].append(tx)
        cumulative_dep = Decimal("0"); cumulative_wd = Decimal("0")
        for tx_date in sorted(by_date):
            day = by_date[tx_date]; cumulative_dep += sum((_dec(t.deposit_amount) for t in day), Decimal("0")); cumulative_wd += sum((_dec(t.withdrawal_amount) for t in day), Decimal("0")); calc = opening + cumulative_dep - cumulative_wd if opening is not None else None
            if account.currency == "KRW":
                max_time = max((t.transaction_time for t in day if t.transaction_time is not None), default=None); latest = [t for t in day if t.transaction_time == max_time] if max_time is not None else day; reported_tx = min(latest, key=lambda t: t.source_row_number)
            else: reported_tx = min(day, key=lambda t: t.source_row_number)
            reported = _dec(reported_tx.balance); diff = None if opening is None else calc - reported; status = "insufficient" if opening is None else "ok" if diff == 0 else "review"
            row = DailyBalanceCheck.query.filter_by(bank_account_id=account.id, check_date=tx_date).first()
            if not row: row = DailyBalanceCheck(bank_account_id=account.id, check_date=tx_date); db.session.add(row)
            row.opening_balance = opening; row.cumulative_deposit = cumulative_dep; row.cumulative_withdrawal = cumulative_wd; row.calculated_balance = calc; row.bank_reported_balance = reported; row.difference = diff; row.status = status; row.checked_at = core.utcnow()


def _post_import():
    try:
        _ensure_annotations(); _refresh_daily_checks(); core.audit("phase2_post_import_validation", user=current_user, target_type="bank_transactions", detail="자동 잔액검증/거래주석 갱신"); db.session.commit()
    except Exception:
        db.session.rollback(); core.app.logger.exception("Phase2 post import validation failed")


def _rate_for(currency, day):
    if currency == "KRW": return Decimal("1")
    row = FxRate.query.filter(FxRate.base_date == day, FxRate.currency == currency, FxRate.round_no == 1).order_by(FxRate.created_at.desc()).first()
    return _dec(row.rate) if row else None


def _status_summary(cutoff, mmt_opening_model):
    krw = {"ok": 0, "review": 0, "insufficient": 0}; fx = {"ok": 0, "review": 0, "insufficient": 0}; details = []
    for a in core.BankAccount.query.filter_by(is_active=True).all():
        if not _account_included(a): continue
        b = _bank_balance(a, cutoff); (krw if a.currency == "KRW" else fx)[b["status"]] += 1; details.append((a, b))
    mmt = {"ok": 0, "review": 0, "insufficient": 0}
    for bank, acct, curr in _all_mmt_keys(mmt_opening_model): mmt[_mmt_balance(mmt_opening_model, bank, acct, curr, cutoff)["status"]] += 1
    ready = krw["review"] == 0 and fx["review"] == 0 and mmt["review"] == 0 and krw["insufficient"] == 0 and fx["insufficient"] == 0 and mmt["insufficient"] == 0
    return {"krw": krw, "fx": fx, "mmt": mmt, "ready": ready, "details": details}


PAGE_STYLE = """<style>.p2head{display:flex;justify-content:space-between;gap:16px;align-items:flex-end;margin-bottom:18px}.p2head h2{margin:0}.muted{color:#667085}.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:12px;margin:14px 0 20px}.kpi{background:#fff;border:1px solid #e4e7ec;border-radius:12px;padding:16px}.kpi b{font-size:23px;display:block;margin-top:5px}.badge{display:inline-block;padding:3px 8px;border-radius:999px;font-size:12px}.badge.ok{background:#ecfdf3;color:#067647}.badge.review{background:#fff4ed;color:#b54708}.badge.insufficient{background:#f2f4f7;color:#475467}.actions{display:flex;gap:8px;flex-wrap:wrap}.actions a,.btnlink{display:inline-block;padding:9px 12px;border-radius:8px;text-decoration:none;background:#155eef;color:#fff}.btnlink.secondary,.actions a.secondary{background:#475467}.filters{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:8px;align-items:end}.num{text-align:right;white-space:nowrap}.small{font-size:12px}.scroll{overflow:auto}.section-title{margin:22px 0 10px}.warnbox{border-left:4px solid #f79009;background:#fffaeb;padding:12px;margin:12px 0}.okbox{border-left:4px solid #12b76a;background:#ecfdf3;padding:12px;margin:12px 0}@media print{.sidebar,.noprint{display:none!important}.content{margin:0!important;padding:0!important}.panel{break-inside:avoid}}</style>"""


def _render(title, body, **ctx):
    tpl = "{% extends 'base.html' %}{% block title %}" + title + "{% endblock %}{% block content %}" + PAGE_STYLE + body + "{% endblock %}"; ctx.setdefault("fmt", _fmt); return render_template_string(tpl, **ctx)


def _snapshot(day, mmt_opening_model):
    accounts = []; krw_total = Decimal("0"); fx_by_currency = defaultdict(Decimal); fx_krw_total = Decimal("0"); missing_rates = []
    for a in core.BankAccount.query.filter_by(is_active=True).all():
        if not _account_included(a): continue
        b = _bank_balance(a, day); bal = b["calculated"]
        if bal is None: continue
        if a.currency == "KRW": krw_total += bal; rate, krw = Decimal("1"), bal
        else:
            fx_by_currency[a.currency] += bal; rate = _rate_for(a.currency, day)
            if rate is None: missing_rates.append(a.currency); krw = None
            else: krw = bal * rate; fx_krw_total += krw
        accounts.append({"bank": a.financial_institution, "account": a.account_number, "name": a.account_name or "", "currency": a.currency, "balance": str(bal), "rate": str(rate) if rate is not None else None, "krw": str(krw) if krw is not None else None, "status": b["status"]})
    mmt_accounts = []; mmt_total = Decimal("0")
    for bank, acct, curr in _all_mmt_keys(mmt_opening_model):
        b = _mmt_balance(mmt_opening_model, bank, acct, curr, day)
        if b["calculated"] is not None and curr == "KRW": mmt_total += b["calculated"]
        mmt_accounts.append({"bank": bank, "account": acct, "currency": curr, "balance": str(b["calculated"]) if b["calculated"] is not None else None, "status": b["status"]})
    loan_by_currency = defaultdict(Decimal)
    for l in core.LoanAccount.query.filter_by(status="active").all(): loan_by_currency[l.currency] += _dec(l.current_balance)
    d1 = day - timedelta(days=1); actual_txs = core.BankTransaction.query.filter(core.BankTransaction.transaction_date == d1).all(); ids = [t.id for t in actual_txs]; ann_map = {x.transaction_id: x for x in TransactionAnnotation.query.filter(TransactionAnnotation.transaction_id.in_(ids) if ids else text("1=0")).all()}
    actual_by_currency = defaultdict(lambda: {"external_in": Decimal("0"), "external_out": Decimal("0"), "internal_in": Decimal("0"), "internal_out": Decimal("0")}); category_totals = defaultdict(Decimal); actual_details = []
    for t in actual_txs:
        ann = ann_map.get(t.id); flow = ann.flow_type if ann else ("external_in" if _dec(t.deposit_amount) > 0 else "external_out"); amount = _dec(t.deposit_amount) if _dec(t.deposit_amount) > 0 else _dec(t.withdrawal_amount); curr = t.account.currency
        if flow in actual_by_currency[curr]: actual_by_currency[curr][flow] += amount
        if ann and ann.category: category_totals[(ann.category.name, curr)] += amount
        actual_details.append({"date": t.transaction_date.isoformat(), "description": t.description or "", "counterparty": t.counterparty or "", "deposit": str(t.deposit_amount), "withdrawal": str(t.withdrawal_amount), "flow": flow, "currency": curr, "category": ann.category.name if ann and ann.category else ""})
    dplus = day + timedelta(days=1); plans = CashPlan.query.filter(CashPlan.planned_date == dplus, CashPlan.status != "cancelled").all(); planned_by_currency = defaultdict(lambda: {"in": Decimal("0"), "out": Decimal("0")}); planned_in_krw = Decimal("0"); planned_out_krw = Decimal("0")
    for p in plans:
        amt = _dec(p.amount); planned_by_currency[p.currency][p.direction] += amt
        if p.currency == "KRW": converted = amt
        else:
            rate = _rate_for(p.currency, day); converted = amt * rate if rate is not None else None
            if rate is None: missing_rates.append(p.currency)
        if converted is not None:
            if p.direction == "in": planned_in_krw += converted
            else: planned_out_krw += converted
    status = _status_summary(day, mmt_opening_model); ready = status["ready"] and not missing_rates; total_financial = krw_total + fx_krw_total + mmt_total
    return {"date": day.isoformat(), "d1": d1.isoformat(), "dplus": dplus.isoformat(), "krw_total": str(krw_total), "fx_krw_total": str(fx_krw_total), "mmt_total": str(mmt_total), "total_financial": str(total_financial), "loan_by_currency": {k: str(v) for k, v in loan_by_currency.items()}, "planned_in": str(planned_in_krw), "planned_out": str(planned_out_krw), "planned_by_currency": {c: {k: str(v) for k, v in vals.items()} for c, vals in planned_by_currency.items()}, "expected": str(total_financial + planned_in_krw - planned_out_krw), "accounts": accounts, "mmt_accounts": mmt_accounts, "actual_by_currency": {c: {k: str(v) for k, v in vals.items()} for c, vals in actual_by_currency.items()}, "category_totals": {f"{name}|{curr}": str(v) for (name, curr), v in category_totals.items()}, "actual_details": actual_details, "plans": [{"direction": p.direction, "counterparty": p.counterparty or "", "description": p.description, "currency": p.currency, "amount": str(p.amount), "status": p.status} for p in plans], "status": {"krw": status["krw"], "fx": status["fx"], "mmt": status["mmt"], "ready": status["ready"]}, "missing_rates": sorted(set(missing_rates)), "ready": ready}


def register():
    global REGISTERED
    if REGISTERED: return
    import app as pkg
    app = pkg.app; mmt_opening_model = pkg.MmtOpeningBalance
    with app.app_context():
        before = {"bank_transactions": core.BankTransaction.query.count(), "opening_balances": core.OpeningBalance.query.count(), "import_batches": core.ImportBatch.query.count(), "import_raw_rows": core.ImportRawRow.query.count()}
        with db.engine.begin() as conn:
            conn.execute(text(f"SELECT pg_advisory_xact_lock({PHASE2_LOCK})")); db.metadata.create_all(bind=conn)
        _ensure_category_seed(); db.session.commit()
        after = {"bank_transactions": core.BankTransaction.query.count(), "opening_balances": core.OpeningBalance.query.count(), "import_batches": core.ImportBatch.query.count(), "import_raw_rows": core.ImportRawRow.query.count()}
        if before != after: raise RuntimeError(f"Phase2 additive migration violated production-data guard: before={before} after={after}")
        _ensure_annotations(); _refresh_daily_checks(); db.session.commit()
    original_confirm = app.view_functions.get("import_confirm")
    if original_confirm and not getattr(original_confirm, "_phase2_wrapped", False):
        def import_confirm_phase2(token):
            response = original_confirm(token)
            try:
                if getattr(response, "status_code", 302) < 400: _post_import()
            except Exception: core.app.logger.exception("Phase2 import post-process wrapper failed")
            return response
        import_confirm_phase2._phase2_wrapped = True; app.view_functions["import_confirm"] = import_confirm_phase2

    @login_required
    def phase2_dashboard():
        day = _latest_transaction_date() or date.today(); snap = _snapshot(day, mmt_opening_model); latest = _latest_upload(); unclassified = TransactionAnnotation.query.filter(TransactionAnnotation.category_id.is_(None)).count()
        body = """<div class="p2head"><div><h2>자금현황 대시보드</h2><p class="muted">기준일 {{ day }}</p></div><div class="actions"><a href="{{ url_for('cash_journal') }}?date={{ day }}">자금일보 보기</a><a class="secondary" href="{{ url_for('phase2_integrity') }}?date={{ day }}">데이터 정합성</a></div></div><div class="kpis"><div class="kpi"><span>원화예금</span><b>{{ fmt(s.krw_total) }}원</b></div><div class="kpi"><span>외화 원화환산</span><b>{{ fmt(s.fx_krw_total) }}원</b></div><div class="kpi"><span>MMT</span><b>{{ fmt(s.mmt_total) }}원</b></div><div class="kpi"><span>총 금융자금</span><b>{{ fmt(s.total_financial) }}원</b></div><div class="kpi"><span>D+1 예정입금</span><b>{{ fmt(s.planned_in) }}원</b></div><div class="kpi"><span>D+1 예정출금</span><b>{{ fmt(s.planned_out) }}원</b></div></div><section class="panel"><h3>데이터 상태</h3><table><tr><th>마지막 e-Branch 업로드</th><td>{{ latest.confirmed_at|kst if latest else '-' }}</td><th>거래자료 최종일</th><td>{{ day }}</td></tr><tr><th>원화계좌</th><td>정상 {{ s.status.krw.ok }} / 확인필요 {{ s.status.krw.review }} / 자료부족 {{ s.status.krw.insufficient }}</td><th>외화계좌</th><td>정상 {{ s.status.fx.ok }} / 확인필요 {{ s.status.fx.review }} / 자료부족 {{ s.status.fx.insufficient }}</td></tr><tr><th>MMT</th><td>정상 {{ s.status.mmt.ok }} / 확인필요 {{ s.status.mmt.review }} / 자료부족 {{ s.status.mmt.insufficient }}</td><th>미분류 거래</th><td>{{ unclassified }}건</td></tr></table>{% if s.ready %}<div class="okbox"><b>자금일보 생성 가능</b></div>{% else %}<div class="warnbox"><b>확인 필요</b> — 데이터 정합성 또는 환율을 확인해 주세요.</div>{% endif %}</section>"""
        return _render("대시보드", body, s=snap, day=day, latest=latest, unclassified=unclassified)

    @login_required
    def phase2_integrity():
        cutoff = _d(request.args.get("date"), _latest_transaction_date() or date.today()); rows=[]
        for a in core.BankAccount.query.order_by(core.BankAccount.financial_institution, core.BankAccount.account_number, core.BankAccount.currency).all(): rows.append({"kind":"예금","account":a,**_bank_balance(a,cutoff)})
        mmt_rows=[]
        for bank,acct,curr in _all_mmt_keys(mmt_opening_model): mmt_rows.append({"bank":bank,"account_no":acct,"currency":curr,**_mmt_balance(mmt_opening_model,bank,acct,curr,cutoff)})
        body="""<div class="p2head"><div><h2>데이터 정합성</h2><p class="muted">2026.01.01 기초잔액 + 누적 입금 - 누적 출금 vs e-Branch 최종잔액</p></div></div><form method="get" class="panel filters noprint"><label>기준일<input type="date" name="date" value="{{ cutoff }}"></label><button>검증 조회</button></form><section class="panel scroll"><h3>예금계좌</h3><table><thead><tr><th>은행</th><th>계좌</th><th>통화</th><th class="num">기초</th><th class="num">누적입금</th><th class="num">누적출금</th><th class="num">계산잔액</th><th class="num">은행잔액</th><th class="num">차이</th><th>상태</th></tr></thead><tbody>{% for r in rows %}<tr><td>{{ r.account.financial_institution }}</td><td>{{ r.account.account_number }}</td><td>{{ r.account.currency }}</td><td class="num">{{ fmt(r.opening,r.account.currency) }}</td><td class="num">{{ fmt(r.deposit,r.account.currency) }}</td><td class="num">{{ fmt(r.withdrawal,r.account.currency) }}</td><td class="num">{{ fmt(r.calculated,r.account.currency) }}</td><td class="num">{{ fmt(r.reported,r.account.currency) }}</td><td class="num">{{ fmt(r.difference,r.account.currency) }}</td><td><span class="badge {{ r.status }}">{{ {'ok':'정상','review':'확인필요','insufficient':'자료부족'}[r.status] }}</span></td></tr>{% endfor %}</tbody></table></section><section class="panel scroll"><h3>MMT</h3><table><thead><tr><th>은행</th><th>계좌</th><th>통화</th><th class="num">기초</th><th class="num">누적입금</th><th class="num">누적출금</th><th class="num">계산잔액</th><th class="num">은행잔액</th><th class="num">차이</th><th>상태</th></tr></thead><tbody>{% for r in mmt_rows %}<tr><td>{{ r.bank }}</td><td>{{ r.account_no }}</td><td>{{ r.currency }}</td><td class="num">{{ fmt(r.opening,r.currency) }}</td><td class="num">{{ fmt(r.deposit,r.currency) }}</td><td class="num">{{ fmt(r.withdrawal,r.currency) }}</td><td class="num">{{ fmt(r.calculated,r.currency) }}</td><td class="num">{{ fmt(r.reported,r.currency) }}</td><td class="num">{{ fmt(r.difference,r.currency) }}</td><td><span class="badge {{ r.status }}">{{ {'ok':'정상','review':'확인필요','insufficient':'자료부족'}[r.status] }}</span></td></tr>{% endfor %}</tbody></table></section>"""
        return _render("데이터 정합성",body,cutoff=cutoff,rows=rows,mmt_rows=mmt_rows)

    @login_required
    def phase2_accounts():
        accounts=core.BankAccount.query.order_by(core.BankAccount.financial_institution,core.BankAccount.account_number,core.BankAccount.currency).all(); settings={s.bank_account_id:s for s in BankAccountSetting.query.all()}
        body="""<div class="p2head"><div><h2>계좌관리</h2><p class="muted">거래가 연결된 금융기관·계좌번호·통화는 변경하지 않습니다.</p></div></div><section class="panel scroll"><table><thead><tr><th>은행</th><th>계좌번호</th><th>통화</th><th>계좌명</th><th>계좌구분</th><th>사용</th><th>자금일보</th><th>내부계좌</th><th>순서</th><th>비고</th><th></th></tr></thead><tbody>{% for a in accounts %}{% set s=settings.get(a.id) %}<tr><form method="post" action="{{ url_for('phase2_account_update',account_id=a.id) }}"><input type="hidden" name="csrf_token" value="{{ csrf_token() }}"><td>{{ a.financial_institution }}</td><td>{{ a.account_number }}</td><td>{{ a.currency }}</td><td><input name="account_name" value="{{ a.account_name or '' }}"></td><td><select name="account_type">{% for x in ['원화 보통예금','외화예금','MMT','기타예금'] %}<option value="{{ x }}" {% if a.account_type==x %}selected{% endif %}>{{ x }}</option>{% endfor %}</select></td><td><input type="checkbox" name="is_active" {% if a.is_active %}checked{% endif %}></td><td><input type="checkbox" name="include" {% if not s or s.include_in_cash_journal %}checked{% endif %}></td><td><input type="checkbox" name="internal" {% if not s or s.is_internal %}checked{% endif %}></td><td><input type="number" name="display_order" value="{{ s.display_order if s else 100 }}" style="width:75px"></td><td><input name="note" value="{{ s.note if s and s.note else '' }}"></td><td><button>저장</button></td></form></tr>{% endfor %}</tbody></table></section>"""
        return _render("계좌관리",body,accounts=accounts,settings=settings)

    @core.roles_required("admin","editor")
    def phase2_account_update(account_id):
        a=db.get_or_404(core.BankAccount,account_id); a.account_name=(request.form.get("account_name") or "").strip(); a.account_type=(request.form.get("account_type") or "기타예금").strip(); a.is_active=request.form.get("is_active")=="on"; s=BankAccountSetting.query.filter_by(bank_account_id=a.id).first()
        if not s: s=BankAccountSetting(bank_account_id=a.id); db.session.add(s)
        s.include_in_cash_journal=request.form.get("include")=="on"; s.is_internal=request.form.get("internal")=="on"; s.display_order=int(request.form.get("display_order") or 100); s.note=(request.form.get("note") or "").strip(); core.audit("bank_account_setting_updated",user=current_user,target_type="bank_account",target_id=a.id,detail=f"name={a.account_name}, type={a.account_type}, active={a.is_active}, include={s.include_in_cash_journal}, internal={s.is_internal}"); db.session.commit(); flash("계좌 관리정보가 저장되었습니다.","success"); return redirect(url_for("phase2_accounts"))

    @login_required
    def phase2_classifications():
        cats=TransactionCategory.query.order_by(TransactionCategory.sort_order,TransactionCategory.name).all(); page=max(1,int(request.args.get("page") or 1)); per=100; total=TransactionAnnotation.query.count(); annotations=TransactionAnnotation.query.join(core.BankTransaction).order_by(core.BankTransaction.transaction_date.desc(),core.BankTransaction.transaction_time.desc().nullslast()).offset((page-1)*per).limit(per).all()
        body="""<div class="p2head"><div><h2>거래분류</h2><p class="muted">은행 Raw Data는 수정하지 않고 관리용 분류만 연결합니다.</p></div></div><section class="panel"><h3>분류 Master — {{ cats|length }}개</h3>{% if current_user.role in ['admin','editor'] %}<form method="post" action="{{ url_for('phase2_category_add') }}" class="filters noprint"><input type="hidden" name="csrf_token" value="{{ csrf_token() }}"><label>분류명<input name="name" required></label><label>표시순서<input type="number" name="sort_order" value="{{ cats|length + 1 }}"></label><button>분류 추가</button></form>{% endif %}<div class="scroll"><table><thead><tr><th>순서</th><th>Code</th><th>분류명</th><th>사용</th><th></th></tr></thead><tbody>{% for c in cats %}<tr><form method="post" action="{{ url_for('phase2_category_update',category_id=c.id) }}"><input type="hidden" name="csrf_token" value="{{ csrf_token() }}"><td><input type="number" name="sort_order" value="{{ c.sort_order }}" style="width:75px"></td><td>{{ c.code }}</td><td><input name="name" value="{{ c.name }}"></td><td><input type="checkbox" name="is_active" {% if c.is_active %}checked{% endif %}></td><td>{% if current_user.role in ['admin','editor'] %}<button>저장</button>{% endif %}</td></form></tr>{% endfor %}</tbody></table></div></section><section class="panel scroll"><h3>거래 분류/내부대체 확인</h3><p class="muted">총 {{ total }}건 · 1페이지 {{ per }}건</p><table><thead><tr><th>일자</th><th>계좌</th><th>적요</th><th>상대방</th><th>입금</th><th>출금</th><th>자동흐름</th><th>분류</th><th>제안</th><th></th></tr></thead><tbody>{% for a in anns %}<tr><form method="post" action="{{ url_for('phase2_annotation_update',annotation_id=a.id) }}"><input type="hidden" name="csrf_token" value="{{ csrf_token() }}"><td>{{ a.transaction.transaction_date }}</td><td>{{ a.transaction.account.account_number }}</td><td>{{ a.transaction.description or '-' }}</td><td>{{ a.transaction.counterparty or '-' }}</td><td class="num">{{ fmt(a.transaction.deposit_amount,a.transaction.account.currency) }}</td><td class="num">{{ fmt(a.transaction.withdrawal_amount,a.transaction.account.currency) }}</td><td><select name="flow_type">{% for f,n in [('external_in','외부입금'),('external_out','외부출금'),('internal_in','내부대체 입금'),('internal_out','내부대체 출금'),('external','기타')] %}<option value="{{ f }}" {% if a.flow_type==f %}selected{% endif %}>{{ n }}</option>{% endfor %}</select></td><td><select name="category_id"><option value="">미분류</option>{% for c in cats if c.is_active %}<option value="{{ c.id }}" {% if a.category_id==c.id %}selected{% endif %}>{{ c.name }}</option>{% endfor %}</select></td><td>{{ a.suggested_category.name if a.suggested_category else '-' }}<div class="small muted">{{ a.suggestion_reason or '' }}</div></td><td><button>확정</button></td></form></tr>{% endfor %}</tbody></table></section><div class="actions noprint">{% if page>1 %}<a class="secondary" href="?page={{ page-1 }}">이전</a>{% endif %}{% if page*per<total %}<a href="?page={{ page+1 }}">다음</a>{% endif %}</div>"""
        return _render("거래분류",body,cats=cats,anns=annotations,total=total,per=per,page=page)

    @core.roles_required("admin","editor")
    def phase2_annotation_update(annotation_id):
        ann=db.get_or_404(TransactionAnnotation,annotation_id); cid=request.form.get("category_id"); ann.category_id=uuid.UUID(cid) if cid else None; ann.flow_type=request.form.get("flow_type") or ann.flow_type; ann.is_confirmed=True; ann.confirmed_by=current_user.id; ann.classification_source="manual"; core.audit("transaction_annotation_confirmed",user=current_user,target_type="bank_transaction",target_id=ann.transaction_id,detail=f"flow={ann.flow_type}, category={cid or 'none'}"); db.session.commit(); return redirect(request.referrer or url_for("phase2_classifications"))

    @core.roles_required("admin","editor")
    def phase2_category_add():
        name=(request.form.get("name") or "").strip()
        if not name: abort(400)
        if TransactionCategory.query.filter(func.lower(TransactionCategory.name)==name.lower()).first(): flash("동일한 거래분류가 이미 존재합니다.","error"); return redirect(url_for("phase2_classifications"))
        max_num=max([int(c.code[1:]) for c in TransactionCategory.query.all() if c.code and c.code.startswith("C") and c.code[1:].isdigit()] or [0]); row=TransactionCategory(code=f"C{max_num+1:03d}",name=name,sort_order=int(request.form.get("sort_order") or TransactionCategory.query.count()+1),is_active=True); db.session.add(row); core.audit("transaction_category_created",user=current_user,target_type="transaction_category",target_id=row.id,detail=name); db.session.commit(); return redirect(url_for("phase2_classifications"))

    @core.roles_required("admin","editor")
    def phase2_category_update(category_id):
        row=db.get_or_404(TransactionCategory,category_id); name=(request.form.get("name") or "").strip(); other=TransactionCategory.query.filter(func.lower(TransactionCategory.name)==name.lower(),TransactionCategory.id!=row.id).first()
        if not name or other: flash("분류명이 비어 있거나 중복되었습니다.","error"); return redirect(url_for("phase2_classifications"))
        old=row.name; row.name=name; row.sort_order=int(request.form.get("sort_order") or row.sort_order); row.is_active=request.form.get("is_active")=="on"; core.audit("transaction_category_updated",user=current_user,target_type="transaction_category",target_id=row.id,detail=f"{old}->{row.name}, active={row.is_active}"); db.session.commit(); return redirect(url_for("phase2_classifications"))

    @login_required
    def phase2_cash_plans():
        selected=_d(request.args.get("date"),(_latest_transaction_date() or date.today())+timedelta(days=1)); plans=CashPlan.query.filter_by(planned_date=selected).order_by(CashPlan.created_at.desc()).all(); accounts=core.BankAccount.query.filter_by(is_active=True).order_by(core.BankAccount.financial_institution,core.BankAccount.account_number).all(); cats=TransactionCategory.query.filter_by(is_active=True).order_by(TransactionCategory.sort_order).all()
        body="""<div class="p2head"><div><h2>자금예정</h2><p class="muted">Phase 2는 manual Source만 사용합니다.</p></div></div><section class="panel noprint"><form method="post" action="{{ url_for('phase2_cash_plan_add') }}" class="filters"><input type="hidden" name="csrf_token" value="{{ csrf_token() }}"><label>예정일<input type="date" name="planned_date" value="{{ selected }}" required></label><label>입/출금<select name="direction"><option value="out">출금</option><option value="in">입금</option></select></label><label>계좌<select name="account_id"><option value="">미지정</option>{% for a in accounts %}<option value="{{ a.id }}">{{ a.financial_institution }} {{ a.account_number }} {{ a.currency }}</option>{% endfor %}</select></label><label>거래처<input name="counterparty"></label><label>내용<input name="description" required></label><label>분류<select name="category_id"><option value="">미분류</option>{% for c in cats %}<option value="{{ c.id }}">{{ c.name }}</option>{% endfor %}</select></label><label>통화<input name="currency" value="KRW" required></label><label>금액<input type="number" step="0.000001" name="amount" required></label><label>상태<select name="status"><option value="planned">예정</option><option value="completed">완료</option><option value="cancelled">취소</option></select></label><label>비고<input name="note"></label><button>등록</button></form></section><form method="get" class="panel filters"><label>조회일<input type="date" name="date" value="{{ selected }}"></label><button>조회</button></form><section class="panel scroll"><table><thead><tr><th>예정일</th><th>입/출금</th><th>계좌</th><th>거래처</th><th>내용</th><th>분류</th><th>통화</th><th class="num">금액</th><th>상태</th><th>Source</th></tr></thead><tbody>{% for p in plans %}<tr><td>{{ p.planned_date }}</td><td>{{ '입금' if p.direction=='in' else '출금' }}</td><td>{{ p.account.account_number if p.account else '-' }}</td><td>{{ p.counterparty or '-' }}</td><td>{{ p.description }}</td><td>{{ p.category.name if p.category else '-' }}</td><td>{{ p.currency }}</td><td class="num">{{ fmt(p.amount,p.currency) }}</td><td>{{ p.status }}</td><td>{{ p.source }}</td></tr>{% else %}<tr><td colspan="10" class="empty">등록된 예정자금이 없습니다.</td></tr>{% endfor %}</tbody></table></section>"""
        return _render("자금예정",body,selected=selected,plans=plans,accounts=accounts,cats=cats)

    @core.roles_required("admin","editor")
    def phase2_cash_plan_add():
        day=_d(request.form.get("planned_date"));
        if not day: abort(400)
        account_id=request.form.get("account_id"); category_id=request.form.get("category_id"); p=CashPlan(planned_date=day,direction=request.form.get("direction") or "out",bank_account_id=uuid.UUID(account_id) if account_id else None,counterparty=(request.form.get("counterparty") or "").strip(),description=(request.form.get("description") or "").strip(),category_id=uuid.UUID(category_id) if category_id else None,currency=(request.form.get("currency") or "KRW").strip().upper(),amount=Decimal(request.form.get("amount") or "0"),status=request.form.get("status") or "planned",note=(request.form.get("note") or "").strip(),source="manual",created_by=current_user.id,updated_by=current_user.id); db.session.add(p); core.audit("cash_plan_created",user=current_user,target_type="cash_plan",target_id=p.id,detail=p.description); db.session.commit(); return redirect(url_for("phase2_cash_plans",date=day.isoformat()))

    @login_required
    def phase2_fx_rates():
        day=_d(request.args.get("date"),_latest_transaction_date() or date.today()); rates=FxRate.query.filter_by(base_date=day).order_by(FxRate.currency).all(); body="""<div class="p2head"><div><h2>환율관리</h2><p class="muted">하나은행 일별 1회차 환율을 사용자 직접등록합니다. 원본 외화금액은 변경하지 않습니다.</p></div></div><section class="panel noprint"><form method="post" action="{{ url_for('phase2_fx_rate_add') }}" class="filters"><input type="hidden" name="csrf_token" value="{{ csrf_token() }}"><label>기준일<input type="date" name="base_date" value="{{ day }}" required></label><label>통화<input name="currency" placeholder="USD" required></label><label>환율<input type="number" step="0.00000001" name="rate" required></label><label>회차<input type="number" name="round_no" value="1" required></label><label>Source<input name="source" value="하나은행" required></label><button>등록/갱신</button></form></section><form method="get" class="panel filters"><label>조회일<input type="date" name="date" value="{{ day }}"></label><button>조회</button></form><section class="panel"><table><thead><tr><th>기준일</th><th>통화</th><th class="num">환율</th><th>회차</th><th>Source</th><th>등록일시</th></tr></thead><tbody>{% for r in rates %}<tr><td>{{ r.base_date }}</td><td>{{ r.currency }}</td><td class="num">{{ r.rate }}</td><td>{{ r.round_no }}</td><td>{{ r.source }}</td><td>{{ r.created_at|kst }}</td></tr>{% else %}<tr><td colspan="6" class="empty">등록된 환율이 없습니다.</td></tr>{% endfor %}</tbody></table></section>"""; return _render("환율관리",body,day=day,rates=rates)

    @core.roles_required("admin","editor")
    def phase2_fx_rate_add():
        day=_d(request.form.get("base_date")); curr=(request.form.get("currency") or "").strip().upper(); round_no=int(request.form.get("round_no") or 1); source=(request.form.get("source") or "하나은행").strip(); row=FxRate.query.filter_by(base_date=day,currency=curr,round_no=round_no,source=source).first()
        if not row: row=FxRate(base_date=day,currency=curr,round_no=round_no,source=source,created_by=current_user.id); db.session.add(row)
        row.rate=Decimal(request.form.get("rate") or "0"); core.audit("fx_rate_saved",user=current_user,target_type="fx_rate",target_id=row.id,detail=f"{day} {curr}={row.rate} {source} {round_no}회차"); db.session.commit(); return redirect(url_for("phase2_fx_rates",date=day.isoformat()))

    @login_required
    def financial_transactions_phase2():
        q=core.BankTransaction.query.join(core.BankAccount); start=_d(request.args.get("start")); end=_d(request.args.get("end")); bank=(request.args.get("bank") or "").strip(); account=(request.args.get("account") or "").strip(); data_type=request.args.get("data_type") or ""; currency=(request.args.get("currency") or "").strip().upper(); direction=request.args.get("direction") or ""; keyword=(request.args.get("q") or "").strip()
        if start:q=q.filter(core.BankTransaction.transaction_date>=start)
        if end:q=q.filter(core.BankTransaction.transaction_date<=end)
        if bank:q=q.filter(core.BankAccount.financial_institution==bank)
        if account:q=q.filter(core.BankAccount.account_number==account)
        if data_type=="krw":q=q.filter(core.BankAccount.currency=="KRW")
        elif data_type=="fx":q=q.filter(core.BankAccount.currency!="KRW")
        if currency:q=q.filter(core.BankAccount.currency==currency)
        if direction=="deposit":q=q.filter(core.BankTransaction.deposit_amount>0)
        elif direction=="withdrawal":q=q.filter(core.BankTransaction.withdrawal_amount>0)
        if keyword:
            like=f"%{keyword}%"; q=q.filter(or_(core.BankTransaction.description.ilike(like),core.BankTransaction.counterparty.ilike(like),core.BankTransaction.detail1.ilike(like),core.BankTransaction.detail2.ilike(like)))
        total=q.count(); per=int(request.args.get("per") or 100); per=per if per in (50,100,200) else 100; page=max(1,int(request.args.get("page") or 1)); rows=q.order_by(core.BankTransaction.transaction_date.desc(),core.BankTransaction.transaction_time.desc().nullslast(),core.BankTransaction.source_row_number.asc()).offset((page-1)*per).limit(per).all(); min_date,max_date=db.session.query(func.min(core.BankTransaction.transaction_date),func.max(core.BankTransaction.transaction_date)).one(); cutoff=end or max_date or date.today(); balances=defaultdict(Decimal)
        for a in core.BankAccount.query.all():
            b=_bank_balance(a,cutoff)
            if b["calculated"] is not None: balances[a.currency]+=b["calculated"]
        banks=[x[0] for x in db.session.query(core.BankAccount.financial_institution).distinct().order_by(core.BankAccount.financial_institution).all()]; accounts=core.BankAccount.query.order_by(core.BankAccount.financial_institution,core.BankAccount.account_number).all(); currencies=[x[0] for x in db.session.query(core.BankAccount.currency).distinct().order_by(core.BankAccount.currency).all()]; args=request.args.to_dict(); from urllib.parse import urlencode; dl_args={k:v for k,v in args.items() if k not in ("page","per")}; qs=urlencode(dl_args); prev_args=dict(args);prev_args["page"]=page-1;prev_args["per"]=per;next_args=dict(args);next_args["page"]=page+1;next_args["per"]=per
        body="""<div class="p2head"><div><h2>금융거래내역</h2><p class="muted">거래자료 {{ min_date or '-' }} ~ {{ max_date or '-' }} · 조회 {{ total }}건</p></div><a class="btnlink" href="{{ url_for('financial_transactions_csv') }}?{{ qs }}">현재 필터 CSV 다운로드</a></div><form class="panel filters" method="get"><label>시작<input type="date" name="start" value="{{ request.args.get('start','') }}"></label><label>종료<input type="date" name="end" value="{{ request.args.get('end','') }}"></label><label>은행<select name="bank"><option value="">전체</option>{% for b in banks %}<option {% if request.args.get('bank')==b %}selected{% endif %}>{{ b }}</option>{% endfor %}</select></label><label>계좌<select name="account"><option value="">전체</option>{% for a in accounts %}<option value="{{ a.account_number }}" {% if request.args.get('account')==a.account_number %}selected{% endif %}>{{ a.account_number }} {{ a.currency }}</option>{% endfor %}</select></label><label>구분<select name="data_type"><option value="">전체</option><option value="krw">원화</option><option value="fx">외화</option></select></label><label>통화<select name="currency"><option value="">전체</option>{% for c in currencies %}<option>{{ c }}</option>{% endfor %}</select></label><label>입출금<select name="direction"><option value="">전체</option><option value="deposit">입금</option><option value="withdrawal">출금</option></select></label><label>검색<input name="q" value="{{ request.args.get('q','') }}"></label><label>건수<select name="per">{% for x in [50,100,200] %}<option value="{{ x }}" {% if per==x %}selected{% endif %}>{{ x }}</option>{% endfor %}</select></label><button>조회</button></form><section class="panel"><div class="actions">{% for c,v in balances.items() %}<span class="badge ok">{{ c }} {{ fmt(v,c) }}</span>{% endfor %}</div></section><section class="panel scroll"><table><thead><tr><th>일자</th><th>시간</th><th>은행</th><th>계좌</th><th>통화</th><th>적요</th><th>상대방</th><th class="num">입금</th><th class="num">출금</th><th class="num">거래후잔액</th></tr></thead><tbody>{% for r in rows %}<tr><td>{{ r.transaction_date }}</td><td>{{ r.transaction_time or '-' }}</td><td>{{ r.account.financial_institution }}</td><td>{{ r.account.account_number }}</td><td>{{ r.account.currency }}</td><td>{{ r.description or '-' }}</td><td>{{ r.counterparty or '-' }}</td><td class="num">{{ fmt(r.deposit_amount,r.account.currency) }}</td><td class="num">{{ fmt(r.withdrawal_amount,r.account.currency) }}</td><td class="num">{{ fmt(r.balance,r.account.currency) }}</td></tr>{% else %}<tr><td colspan="10" class="empty">조회결과가 없습니다.</td></tr>{% endfor %}</tbody></table></section><div class="actions">{% if page>1 %}<a class="secondary" href="?{{ prev_qs }}">이전</a>{% endif %}{% if page*per<total %}<a href="?{{ next_qs }}">다음</a>{% endif %}</div>"""
        return _render("금융거래내역",body,rows=rows,total=total,per=per,page=page,min_date=min_date,max_date=max_date,balances=balances,banks=banks,accounts=accounts,currencies=currencies,qs=qs,prev_qs=urlencode(prev_args),next_qs=urlencode(next_args),request=request)

    @login_required
    def financial_transactions_csv_phase2():
        q=core.BankTransaction.query.join(core.BankAccount); start=_d(request.args.get("start")); end=_d(request.args.get("end")); bank=(request.args.get("bank") or "").strip(); account=(request.args.get("account") or "").strip(); data_type=request.args.get("data_type") or ""; currency=(request.args.get("currency") or "").strip().upper(); direction=request.args.get("direction") or ""; keyword=(request.args.get("q") or "").strip()
        if start:q=q.filter(core.BankTransaction.transaction_date>=start)
        if end:q=q.filter(core.BankTransaction.transaction_date<=end)
        if bank:q=q.filter(core.BankAccount.financial_institution==bank)
        if account:q=q.filter(core.BankAccount.account_number==account)
        if data_type=="krw":q=q.filter(core.BankAccount.currency=="KRW")
        elif data_type=="fx":q=q.filter(core.BankAccount.currency!="KRW")
        if currency:q=q.filter(core.BankAccount.currency==currency)
        if direction=="deposit":q=q.filter(core.BankTransaction.deposit_amount>0)
        elif direction=="withdrawal":q=q.filter(core.BankTransaction.withdrawal_amount>0)
        if keyword:
            like=f"%{keyword}%";q=q.filter(or_(core.BankTransaction.description.ilike(like),core.BankTransaction.counterparty.ilike(like),core.BankTransaction.detail1.ilike(like),core.BankTransaction.detail2.ilike(like)))
        rows=q.order_by(core.BankTransaction.transaction_date,core.BankTransaction.transaction_time,core.BankTransaction.source_row_number).all();out=io.StringIO();import csv;w=csv.writer(out);w.writerow(["일자","시간","은행","계좌","통화","적요","상대방","입금","출금","잔액","Import Batch"])
        for r in rows:w.writerow([r.transaction_date,r.transaction_time or "",r.account.financial_institution,r.account.account_number,r.account.currency,r.description or "",r.counterparty or "",r.deposit_amount,r.withdrawal_amount,r.balance,r.batch.batch_code])
        return Response("\ufeff"+out.getvalue(),mimetype="text/csv; charset=utf-8",headers={"Content-Disposition":"attachment; filename=bank_transactions_filtered.csv"})

    @login_required
    def mmt_phase2():
        cutoff=_d(request.args.get("date"),_latest_transaction_date() or date.today());rows=[];total=Decimal("0");counts={"ok":0,"review":0,"insufficient":0}
        for bank,acct,curr in _all_mmt_keys(mmt_opening_model):
            b=_mmt_balance(mmt_opening_model,bank,acct,curr,cutoff);rows.append({"bank":bank,"account":acct,"currency":curr,**b});counts[b["status"]]+=1
            if curr=="KRW" and b["calculated"] is not None:total+=b["calculated"]
        body="""<div class="p2head"><div><h2>MMT</h2><p class="muted">일반 보통예금과 분리하여 계좌별 검증합니다.</p></div></div><form class="panel filters" method="get"><label>기준일<input type="date" name="date" value="{{ cutoff }}"></label><button>조회</button></form><div class="kpis"><div class="kpi"><span>MMT 계좌수</span><b>{{ rows|length }}</b></div><div class="kpi"><span>기준일 총잔액(KRW)</span><b>{{ fmt(total) }}원</b></div><div class="kpi"><span>정상</span><b>{{ counts.ok }}</b></div><div class="kpi"><span>확인필요/자료부족</span><b>{{ counts.review+counts.insufficient }}</b></div></div><section class="panel scroll"><table><thead><tr><th>은행</th><th>계좌</th><th>통화</th><th class="num">기초</th><th class="num">입금</th><th class="num">출금</th><th class="num">계산잔액</th><th class="num">e-Branch 잔액</th><th class="num">차이</th><th>상태</th></tr></thead><tbody>{% for r in rows %}<tr><td>{{ r.bank }}</td><td>{{ r.account }}</td><td>{{ r.currency }}</td><td class="num">{{ fmt(r.opening,r.currency) }}</td><td class="num">{{ fmt(r.deposit,r.currency) }}</td><td class="num">{{ fmt(r.withdrawal,r.currency) }}</td><td class="num">{{ fmt(r.calculated,r.currency) }}</td><td class="num">{{ fmt(r.reported,r.currency) }}</td><td class="num">{{ fmt(r.difference,r.currency) }}</td><td><span class="badge {{ r.status }}">{{ {'ok':'정상','review':'확인필요','insufficient':'자료부족'}[r.status] }}</span></td></tr>{% endfor %}</tbody></table></section>""";return _render("MMT",body,cutoff=cutoff,rows=rows,total=total,counts=counts)

    @login_required
    def cash_journal_phase2():
        day=_d(request.args.get("date"),_latest_transaction_date() or date.today());snap=_snapshot(day,mmt_opening_model);versions=CashJournalVersion.query.filter_by(journal_date=day).order_by(CashJournalVersion.version.desc()).all()
        body="""<div class="p2head"><div><h2>자금일보</h2><p class="muted">D-1 실적 · D-Day 잔액 · D+1 예정</p></div><div class="actions noprint"><a href="{{ url_for('phase2_cash_journal_excel') }}?date={{ day }}">Excel 다운로드</a><a class="secondary" href="javascript:window.print()">인쇄</a></div></div><form method="get" class="panel filters noprint"><label>기준일<input type="date" name="date" value="{{ day }}"></label><button>조회</button></form><div class="kpis"><div class="kpi"><span>원화예금</span><b>{{ fmt(s.krw_total) }}원</b></div><div class="kpi"><span>외화 환산</span><b>{{ fmt(s.fx_krw_total) }}원</b></div><div class="kpi"><span>MMT</span><b>{{ fmt(s.mmt_total) }}원</b></div><div class="kpi"><span>총 금융자금</span><b>{{ fmt(s.total_financial) }}원</b></div><div class="kpi"><span>D+1 예상자금</span><b>{{ fmt(s.expected) }}원</b></div></div>{% if s.ready %}<div class="okbox">데이터 상태: <b>자금일보 확정 가능</b></div>{% else %}<div class="warnbox">데이터 상태: <b>확정 불가</b> — 잔액 불일치/자료부족/환율 누락을 확인하세요.{% if s.missing_rates %} 환율누락: {{ s.missing_rates|join(', ') }}{% endif %}</div>{% endif %}<h3 class="section-title">D-1 실적 — {{ s.d1 }}</h3><section class="panel"><table><thead><tr><th>통화</th><th class="num">외부입금</th><th class="num">외부출금</th><th class="num">내부대체 입금</th><th class="num">내부대체 출금</th></tr></thead><tbody>{% for c,v in s.actual_by_currency.items() %}<tr><td>{{ c }}</td><td class="num">{{ fmt(v.external_in,c) }}</td><td class="num">{{ fmt(v.external_out,c) }}</td><td class="num">{{ fmt(v.internal_in,c) }}</td><td class="num">{{ fmt(v.internal_out,c) }}</td></tr>{% else %}<tr><td colspan="5" class="empty">거래 없음</td></tr>{% endfor %}</tbody></table><div class="actions">{% for key,v in s.category_totals.items() %}{% set parts=key.split('|') %}<span class="badge ok">{{ parts[0] }} / {{ parts[1] }} {{ fmt(v,parts[1]) }}</span>{% endfor %}</div></section><h3 class="section-title">D-Day 계좌별 자금잔액 — {{ s.date }}</h3><section class="panel scroll"><table><thead><tr><th>은행</th><th>계좌</th><th>통화</th><th class="num">원통화잔액</th><th class="num">적용환율</th><th class="num">원화환산</th><th>상태</th></tr></thead><tbody>{% for a in s.accounts %}<tr><td>{{ a.bank }}</td><td>{{ a.account }}</td><td>{{ a.currency }}</td><td class="num">{{ a.balance }}</td><td class="num">{{ a.rate or '-' }}</td><td class="num">{{ a.krw or '-' }}</td><td><span class="badge {{ a.status }}">{{ {'ok':'정상','review':'확인필요','insufficient':'자료부족'}[a.status] }}</span></td></tr>{% endfor %}</tbody></table></section><h3 class="section-title">D+1 예정 — {{ s.dplus }}</h3><section class="panel"><div class="kpis"><div class="kpi"><span>예정입금(KRW환산)</span><b>{{ fmt(s.planned_in) }}</b></div><div class="kpi"><span>예정출금(KRW환산)</span><b>{{ fmt(s.planned_out) }}</b></div></div><table><thead><tr><th>입/출금</th><th>거래처</th><th>내용</th><th>통화</th><th class="num">금액</th><th>상태</th></tr></thead><tbody>{% for p in s.plans %}<tr><td>{{ '입금' if p.direction=='in' else '출금' }}</td><td>{{ p.counterparty }}</td><td>{{ p.description }}</td><td>{{ p.currency }}</td><td class="num">{{ p.amount }}</td><td>{{ p.status }}</td></tr>{% else %}<tr><td colspan="6" class="empty">예정내역 없음</td></tr>{% endfor %}</tbody></table></section><section class="panel noprint"><h3>상태 / Version</h3><form method="post" action="{{ url_for('phase2_cash_journal_action') }}" class="filters"><input type="hidden" name="csrf_token" value="{{ csrf_token() }}"><input type="hidden" name="date" value="{{ day }}"><label>수정사유/메모<input name="reason"></label><button name="action" value="review">검토완료 Snapshot</button><button name="action" value="confirm" {% if not s.ready %}disabled{% endif %}>확정 Snapshot</button></form><p class="muted">기존 Version은 삭제하지 않습니다.</p><div class="actions">{% for v in versions %}<span class="badge {{ 'ok' if v.status=='confirmed' else 'review' }}">V{{ v.version }} {{ v.status }} {{ v.created_at|kst }}</span>{% endfor %}</div></section>"""
        return _render("자금일보",body,day=day,s=snap,versions=versions)

    @core.roles_required("admin","editor")
    def phase2_cash_journal_action():
        day=_d(request.form.get("date"));action=request.form.get("action")
        if action not in ("review","confirm") or not day:abort(400)
        snap=_snapshot(day,mmt_opening_model)
        if action=="confirm" and not snap["ready"]:flash("잔액 불일치·자료부족·환율 누락이 있어 확정할 수 없습니다.","error");return redirect(url_for("cash_journal",date=day.isoformat()))
        next_ver=(db.session.query(func.max(CashJournalVersion.version)).filter_by(journal_date=day).scalar() or 0)+1;status="confirmed" if action=="confirm" else "reviewed";v=CashJournalVersion(journal_date=day,version=next_ver,status=status,snapshot_json=snap,created_by=current_user.id,reviewed_by=current_user.id if action=="review" else None,confirmed_by=current_user.id if action=="confirm" else None,confirmed_at=core.utcnow() if action=="confirm" else None,change_reason=(request.form.get("reason") or "").strip());db.session.add(v);core.audit("cash_journal_snapshot_created",user=current_user,target_type="cash_journal_version",target_id=v.id,detail=f"date={day} version={next_ver} status={status}");db.session.commit();flash(f"자금일보 Version {next_ver} ({status}) 저장 완료","success");return redirect(url_for("cash_journal",date=day.isoformat()))

    @login_required
    def phase2_cash_journal_excel():
        day=_d(request.args.get("date"),_latest_transaction_date() or date.today());snap=_snapshot(day,mmt_opening_model);wb=Workbook();ws=wb.active;ws.title="자금일보";ws.append(["MedPark 자금일보",day.isoformat()]);ws["A1"].font=Font(bold=True,size=16);ws.append(["원화예금",Decimal(snap["krw_total"])]);ws.append(["외화 원화환산",Decimal(snap["fx_krw_total"])]);ws.append(["MMT",Decimal(snap["mmt_total"])]);ws.append(["총 금융자금",Decimal(snap["total_financial"])]);ws.append(["D+1 예정입금",Decimal(snap["planned_in"])]);ws.append(["D+1 예정출금",Decimal(snap["planned_out"])]);ws.append(["D+1 예상자금",Decimal(snap["expected"])]);ws.append([]);ws.append(["은행","계좌","통화","원통화잔액","적용환율","원화환산","상태"])
        for a in snap["accounts"]:ws.append([a["bank"],a["account"],a["currency"],Decimal(a["balance"]),Decimal(a["rate"]) if a["rate"] else None,Decimal(a["krw"]) if a["krw"] else None,a["status"]])
        ws.append([]);ws.append(["D+1","입/출금","거래처","내용","통화","금액","상태"])
        for p in snap["plans"]:ws.append([snap["dplus"],"입금" if p["direction"]=="in" else "출금",p["counterparty"],p["description"],p["currency"],Decimal(p["amount"]),p["status"]])
        for col in "ABCDEFG":ws.column_dimensions[col].width=18
        out=io.BytesIO();wb.save(out);out.seek(0);return Response(out.getvalue(),mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f"attachment; filename=MedPark_CashJournal_{day.isoformat()}.xlsx"})

    app.add_url_rule("/phase2/integrity","phase2_integrity",phase2_integrity,methods=["GET"])
    app.add_url_rule("/reference/accounts","phase2_accounts",phase2_accounts,methods=["GET"])
    app.add_url_rule("/reference/accounts/<uuid:account_id>","phase2_account_update",phase2_account_update,methods=["POST"])
    app.add_url_rule("/reference/classifications","phase2_classifications",phase2_classifications,methods=["GET"])
    app.add_url_rule("/reference/classifications/add","phase2_category_add",phase2_category_add,methods=["POST"])
    app.add_url_rule("/reference/classifications/category/<uuid:category_id>","phase2_category_update",phase2_category_update,methods=["POST"])
    app.add_url_rule("/reference/classifications/<uuid:annotation_id>","phase2_annotation_update",phase2_annotation_update,methods=["POST"])
    app.add_url_rule("/cash-plans","phase2_cash_plans",phase2_cash_plans,methods=["GET"])
    app.add_url_rule("/cash-plans/add","phase2_cash_plan_add",phase2_cash_plan_add,methods=["POST"])
    app.add_url_rule("/reference/fx-rates","phase2_fx_rates",phase2_fx_rates,methods=["GET"])
    app.add_url_rule("/reference/fx-rates/add","phase2_fx_rate_add",phase2_fx_rate_add,methods=["POST"])
    app.add_url_rule("/cash-journal/action","phase2_cash_journal_action",phase2_cash_journal_action,methods=["POST"])
    app.add_url_rule("/cash-journal.xlsx","phase2_cash_journal_excel",phase2_cash_journal_excel,methods=["GET"])
    app.view_functions["dashboard"]=phase2_dashboard;app.view_functions["financial_transactions"]=financial_transactions_phase2;app.view_functions["financial_transactions_csv"]=financial_transactions_csv_phase2;app.view_functions["mmt"]=mmt_phase2;app.view_functions["cash_journal"]=cash_journal_phase2;REGISTERED=True
