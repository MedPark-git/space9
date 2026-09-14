"""MedPark Cash entry package.

This package intentionally shadows the legacy root app.py so the production
start command remains `app:app` while phase-1 opening-balance XLSX features
are layered on top without rewriting the stable core application.
"""
import hashlib
import importlib.util
import io
import json
import shutil
import sys
import uuid
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from flask import Response, abort, flash, redirect, render_template, request, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook as XlsxWorkbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

_core_path = Path(__file__).resolve().parent.parent / "app.py"
_spec = importlib.util.spec_from_file_location("medpark_cash_core", _core_path)
core = importlib.util.module_from_spec(_spec)
sys.modules[_spec.name] = core
_spec.loader.exec_module(core)

app = core.app
db = core.db

OPENING_HEADERS = ["구분", "금융기관", "계좌번호", "계좌명", "통화", "기준일", "기초잔액", "비고"]
OPENING_BASE_DATE = date(2026, 1, 1)


class MmtOpeningBalance(db.Model):
    __tablename__ = "mmt_opening_balances"
    id = db.Column(db.Uuid(as_uuid=True), primary_key=True, default=uuid.uuid4)
    base_date = db.Column(db.Date, nullable=False, index=True)
    financial_institution = db.Column(db.String(80), nullable=False)
    account_number = db.Column(db.String(100), nullable=False)
    account_name = db.Column(db.String(200))
    currency = db.Column(db.String(10), nullable=False, default="KRW")
    opening_balance = db.Column(db.Numeric(28, 6), nullable=False)
    status = db.Column(db.String(30), nullable=False, default="confirmed_manual")
    calculation_method = db.Column(db.String(500), nullable=False, default="기초잔액 Excel 일괄등록")
    source_batch_id = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("import_batches.id"))
    created_by = db.Column(db.Uuid(as_uuid=True), db.ForeignKey("users.id"), nullable=False)
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow)
    updated_at = db.Column(db.DateTime(timezone=True), nullable=False, default=core.utcnow, onupdate=core.utcnow)
    source_batch = db.relationship("ImportBatch")
    __table_args__ = (
        db.UniqueConstraint("base_date", "financial_institution", "account_number", "currency", name="uq_mmt_opening_identity"),
    )


def _parse_opening_balance_xlsx(path):
    try:
        wb = load_workbook(path, read_only=False, data_only=False)
    except Exception as exc:
        raise ValueError("정상적인 .xlsx 기초잔액 양식이 아닙니다.") from exc
    if "기초잔액입력" not in wb.sheetnames:
        raise ValueError("[기초잔액입력] 시트를 찾을 수 없습니다. 제공된 양식을 사용해 주세요.")
    ws = wb["기초잔액입력"]
    header_row = None
    for r in range(1, min(ws.max_row, 20) + 1):
        vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, 9)]
        if vals == OPENING_HEADERS:
            header_row = r
            break
    if not header_row:
        raise ValueError("기초잔액 양식 Header가 변경되었습니다. 제공된 양식을 다시 내려받아 작성해 주세요.")

    rows, errors, seen, totals = [], [], set(), {}
    for r in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, 9)]
        if all(v in (None, "") for v in values):
            continue
        row = dict(zip(OPENING_HEADERS, values))
        kind = str(row["구분"] or "").strip().upper()
        kind = "MMT" if kind == "MMT" else "예금" if kind in {"예금", "BANK", "DEPOSIT"} else kind
        bank = str(row["금융기관"] or "").strip()
        account = str(row["계좌번호"] or "").strip()
        account_name = str(row["계좌명"] or "").strip()
        currency = str(row["통화"] or "").strip().upper()
        note = str(row["비고"] or "").strip()

        if ws.cell(r, 7).data_type == "f":
            errors.append({"row": r, "message": "기초잔액은 수식이 아닌 확정 금액을 직접 입력해야 합니다."})
            continue
        try:
            raw_date = row["기준일"]
            if isinstance(raw_date, datetime):
                base_date = raw_date.date()
            elif isinstance(raw_date, date):
                base_date = raw_date
            else:
                base_date = datetime.strptime(str(raw_date or "").strip()[:10], "%Y-%m-%d").date()
        except Exception:
            errors.append({"row": r, "message": "기준일은 2026-01-01 형식이어야 합니다."})
            continue
        if kind not in {"예금", "MMT"}:
            errors.append({"row": r, "message": "구분은 예금 또는 MMT만 사용할 수 있습니다."})
            continue
        if not bank or not account or not currency:
            errors.append({"row": r, "message": "금융기관·계좌번호·통화는 필수입니다."})
            continue
        if base_date != OPENING_BASE_DATE:
            errors.append({"row": r, "message": "이번 등록 기준일은 2026-01-01만 허용합니다."})
            continue
        if row["기초잔액"] in (None, ""):
            errors.append({"row": r, "message": "기초잔액이 비어 있습니다. 잔액이 없으면 0을 입력하세요."})
            continue
        try:
            balance = Decimal(str(row["기초잔액"]).replace(",", ""))
        except Exception:
            errors.append({"row": r, "message": "기초잔액은 숫자로 입력해야 합니다."})
            continue
        key = (kind, bank, account, currency)
        if key in seen:
            errors.append({"row": r, "message": "동일한 구분·금융기관·계좌번호·통화가 파일 내 중복되었습니다."})
            continue
        seen.add(key)
        totals[currency] = totals.get(currency, Decimal("0")) + balance
        rows.append({
            "source_row_number": r, "kind": kind, "bank": bank, "account_number": account,
            "account_name": account_name, "currency": currency, "base_date": base_date.isoformat(),
            "opening_balance": str(balance), "note": note,
        })
    if not rows and not errors:
        raise ValueError("등록할 기초잔액 행이 없습니다.")
    structure_hash = hashlib.sha256("|".join(OPENING_HEADERS).encode("utf-8")).hexdigest()
    return {"rows": rows, "errors": errors, "structure_hash": structure_hash, "totals": {k: str(v) for k, v in totals.items()}}


def _action(row):
    base_date = core.parse_iso_date(row["base_date"])
    if row["kind"] == "MMT":
        existing = MmtOpeningBalance.query.filter_by(
            base_date=base_date, financial_institution=row["bank"], account_number=row["account_number"], currency=row["currency"]
        ).first()
        return "update" if existing else "new"
    account = core.BankAccount.query.filter_by(
        financial_institution=row["bank"], account_number=row["account_number"], currency=row["currency"]
    ).first()
    if not account:
        return "new_account"
    existing = core.OpeningBalance.query.filter_by(base_date=base_date, bank_account_id=account.id).first()
    return "update" if existing else "new"


def _template_rows():
    rows = []
    for a in core.BankAccount.query.order_by(core.BankAccount.financial_institution, core.BankAccount.account_number, core.BankAccount.currency).all():
        rows.append(["예금", a.financial_institution, a.account_number, a.account_name or "", a.currency, OPENING_BASE_DATE, None, ""])
    seen = set()
    for t in core.MmtTransaction.query.order_by(core.MmtTransaction.financial_institution, core.MmtTransaction.account_number).all():
        key = (t.financial_institution, t.account_number, t.currency)
        if key not in seen:
            rows.append(["MMT", t.financial_institution, t.account_number, t.account_name or "", t.currency, OPENING_BASE_DATE, None, ""])
            seen.add(key)
    for t in MmtOpeningBalance.query.filter_by(base_date=OPENING_BASE_DATE).order_by(MmtOpeningBalance.financial_institution, MmtOpeningBalance.account_number).all():
        key = (t.financial_institution, t.account_number, t.currency)
        if key not in seen:
            rows.append(["MMT", t.financial_institution, t.account_number, t.account_name or "", t.currency, OPENING_BASE_DATE, None, ""])
            seen.add(key)
    return rows


def _build_template():
    wb = XlsxWorkbook(); ws = wb.active; ws.title = "기초잔액입력"
    ws.merge_cells("A1:H1"); ws["A1"] = "MedPark 2026년 기초잔액 일괄등록 양식"
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=15); ws["A1"].fill = PatternFill("solid", fgColor="173B57"); ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:H2"); ws["A2"] = "G열 기초잔액을 모두 입력하세요. 0원 계좌도 빈칸이 아닌 0을 입력해야 합니다."
    ws["A2"].fill = PatternFill("solid", fgColor="EAF2F8"); ws["A2"].font = Font(bold=True, color="173B57")
    for c, h in enumerate(OPENING_HEADERS, 1):
        cell = ws.cell(5, c, h); cell.font = Font(bold=True, color="173B57"); cell.fill = PatternFill("solid", fgColor="D9E6F2"); cell.alignment = Alignment(horizontal="center")
    for r_idx, row in enumerate(_template_rows(), 6):
        for c_idx, val in enumerate(row, 1): ws.cell(r_idx, c_idx, val)
        ws.cell(r_idx, 6).number_format = "yyyy-mm-dd"; ws.cell(r_idx, 7).number_format = "#,##0.######"; ws.cell(r_idx, 7).fill = PatternFill("solid", fgColor="FFF7D6")
    for col, width in {"A":10,"B":11,"C":24,"D":31,"E":10,"F":14,"G":20,"H":32}.items(): ws.column_dimensions[col].width = width
    ws.freeze_panes = "A6"
    guide = wb.create_sheet("작성안내")
    guide.append(["작성 원칙", "내용"]); guide.append(["기초잔액", "실제 확정금액을 직접 입력합니다. 0원도 0 입력."]); guide.append(["기준일", "2026-01-01 고정"]); guide.append(["통화", "원통화 그대로 입력하며 환율환산하지 않습니다."]); guide.append(["확정 우선순위", "Excel로 등록한 기초잔액은 e-Branch 자동산출이 덮어쓰지 않습니다."])
    out = io.BytesIO(); wb.save(out); out.seek(0); return out


def opening_balances_view():
    base_date = OPENING_BASE_DATE
    rows = core.OpeningBalance.query.filter_by(base_date=base_date).join(core.BankAccount).order_by(
        core.BankAccount.financial_institution, core.BankAccount.account_number, core.BankAccount.currency
    ).all()
    mmt_rows = MmtOpeningBalance.query.filter_by(base_date=base_date).order_by(
        MmtOpeningBalance.financial_institution, MmtOpeningBalance.account_number, MmtOpeningBalance.currency
    ).all()
    manual_count = sum(1 for r in rows if r.status == "confirmed_manual") + len(mmt_rows)
    auto_count = sum(1 for r in rows if r.status == "confirmed")
    review_count = sum(1 for r in rows if r.status == "needs_review")
    return render_template("opening_balances.html", rows=rows, mmt_rows=mmt_rows, base_date=base_date, manual_count=manual_count, auto_count=auto_count, review_count=review_count)


def opening_balance_template_view():
    out = _build_template()
    return Response(out.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=medpark_opening_balances_2026.xlsx"})


def opening_balance_upload_view():
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename or Path(uploaded.filename).suffix.lower() != ".xlsx":
        flash("기초잔액 등록은 제공된 .xlsx 양식만 사용할 수 있습니다.", "error")
        return redirect(url_for("opening_balances"))
    token = uuid.uuid4().hex; xlsx_path = core.PENDING_DIR / f"{token}_opening.xlsx"; uploaded.save(xlsx_path)
    try:
        parsed = _parse_opening_balance_xlsx(xlsx_path)
        preview_rows = []
        counts = {"total": len(parsed["rows"]), "new": 0, "update": 0, "new_account": 0, "error": len(parsed["errors"]), "bank": 0, "mmt": 0}
        for row in parsed["rows"]:
            action = _action(row); item = dict(row); item["action"] = action; preview_rows.append(item); counts[action] += 1; counts["mmt" if row["kind"] == "MMT" else "bank"] += 1
        preview = {"token": token, "kind": "opening", "created_by": str(current_user.id), "original_filename": uploaded.filename, "file_sha256": core.file_sha256(xlsx_path), "structure_hash": parsed["structure_hash"], "rows": preview_rows, "errors": parsed["errors"], "totals": parsed["totals"], "counts": counts}
        with core.preview_path(token).open("w", encoding="utf-8") as f: json.dump(preview, f, ensure_ascii=False)
        core.audit("opening_balance_preview_created", user=current_user, target_type="opening_balance_import", target_id=token, detail=f"rows={counts['total']} errors={counts['error']}"); core.commit_or_rollback()
        return redirect(url_for("opening_balance_preview", token=token))
    except Exception as exc:
        xlsx_path.unlink(missing_ok=True); core.app.logger.exception("Opening balance preview failed"); flash(str(exc), "error"); return redirect(url_for("opening_balances"))


def opening_balance_preview_view(token):
    data = core.load_preview(token)
    if data.get("kind") != "opening": abort(404)
    return render_template("opening_balance_preview.html", data=data)


def opening_balance_confirm_view(token):
    data = core.load_preview(token); xlsx_path = core.PENDING_DIR / f"{token}_opening.xlsx"
    if data.get("kind") != "opening" or not xlsx_path.exists() or core.file_sha256(xlsx_path) != data.get("file_sha256"):
        flash("기초잔액 원본 파일 검증에 실패했습니다. 다시 업로드해 주세요.", "error"); return redirect(url_for("opening_balances"))
    try:
        parsed = _parse_opening_balance_xlsx(xlsx_path)
        if parsed["errors"]:
            flash("오류가 남아 있어 등록 확정할 수 없습니다. Excel을 수정해 다시 업로드해 주세요.", "error"); return redirect(url_for("opening_balance_preview", token=token))
        actions = [_action(r) for r in parsed["rows"]]
        batch_code = core.next_batch_code(); stored_file = core.IMPORT_DIR / f"{batch_code}_opening.xlsx"; shutil.copy2(xlsx_path, stored_file)
        batch = core.ImportBatch(batch_code=batch_code, file_type="opening", original_filename=data["original_filename"], stored_path=str(stored_file), file_sha256=data["file_sha256"], structure_hash=parsed["structure_hash"], query_start_date=OPENING_BASE_DATE, query_end_date=OPENING_BASE_DATE, total_rows=len(parsed["rows"]), new_rows=sum(1 for a in actions if a in {"new", "new_account"}), duplicate_rows=sum(1 for a in actions if a == "update"), review_rows=0, error_rows=0, status="confirmed", created_by=current_user.id, confirmed_at=core.utcnow())
        db.session.add(batch); db.session.flush(); bank_cache = {}
        for row, action in zip(parsed["rows"], actions):
            raw = {"구분": row["kind"], "금융기관": row["bank"], "계좌번호": row["account_number"], "계좌명": row["account_name"], "통화": row["currency"], "기준일": row["base_date"], "기초잔액": row["opening_balance"], "비고": row["note"]}
            db.session.add(core.ImportRawRow(import_batch_id=batch.id, source_row_number=row["source_row_number"], row_status="updated" if action == "update" else "new", message="기초잔액 Excel 확정", raw_data=raw))
            bal = core.dec(row["opening_balance"]); base_date = core.parse_iso_date(row["base_date"])
            if row["kind"] == "MMT":
                ob = MmtOpeningBalance.query.filter_by(base_date=base_date, financial_institution=row["bank"], account_number=row["account_number"], currency=row["currency"]).first()
                if not ob:
                    ob = MmtOpeningBalance(base_date=base_date, financial_institution=row["bank"], account_number=row["account_number"], account_name=row["account_name"], currency=row["currency"], opening_balance=bal, created_by=current_user.id); db.session.add(ob)
                ob.account_name = row["account_name"] or ob.account_name; ob.opening_balance = bal; ob.status = "confirmed_manual"; ob.calculation_method = "기초잔액 Excel 일괄등록(사용자 확정)"; ob.source_batch_id = batch.id
            else:
                account = core.get_or_create_bank_account(bank_cache, {"bank": row["bank"], "account_number": row["account_number"], "currency": row["currency"], "account_name": row["account_name"]})
                if row["account_name"]: account.account_name = row["account_name"]
                ob = core.OpeningBalance.query.filter_by(base_date=base_date, bank_account_id=account.id).first()
                if not ob:
                    ob = core.OpeningBalance(base_date=base_date, bank_account_id=account.id, currency=row["currency"], created_by=current_user.id); db.session.add(ob)
                ob.source_date = None; ob.source_balance = None; ob.opening_balance = bal; ob.status = "confirmed_manual"; ob.calculation_method = "기초잔액 Excel 일괄등록(사용자 확정)"; ob.source_batch_id = batch.id; ob.source_transaction_id = None
        core.audit("opening_balance_excel_confirmed", user=current_user, target_type="import_batch", target_id=batch.id, detail=f"{batch_code} rows={batch.total_rows} updated={batch.duplicate_rows}"); db.session.commit()
        xlsx_path.unlink(missing_ok=True); core.preview_path(token).unlink(missing_ok=True)
        flash(f"기초잔액 등록 완료: 총 {batch.total_rows}계좌 · 신규 {batch.new_rows} · 기존 갱신 {batch.duplicate_rows}", "success"); return redirect(url_for("opening_balances"))
    except Exception:
        db.session.rollback(); core.app.logger.exception("Opening balance confirmation failed"); flash("기초잔액 등록 확정 중 오류가 발생했습니다. DB에는 반영되지 않았습니다.", "error"); return redirect(url_for("opening_balance_preview", token=token))


def generate_opening_balances_view():
    base_date = OPENING_BASE_DATE; source_date = date(2025, 12, 31)
    accounts = core.BankAccount.query.order_by(core.BankAccount.financial_institution, core.BankAccount.account_number, core.BankAccount.currency).all()
    confirmed = review = manual_kept = 0
    for account in accounts:
        ob = core.OpeningBalance.query.filter_by(base_date=base_date, bank_account_id=account.id).first()
        if ob and ob.status == "confirmed_manual":
            manual_kept += 1; continue
        txs = core.BankTransaction.query.filter_by(account_id=account.id, transaction_date=source_date).order_by(core.BankTransaction.source_row_number).all()
        helper_rows = [{"id":str(t.id), "transaction_time":t.transaction_time.isoformat() if t.transaction_time else None, "balance":str(t.balance), "deposit":str(t.deposit_amount), "withdrawal":str(t.withdrawal_amount)} for t in txs]
        final_data, _status_label, method = core.select_final_balance(helper_rows)
        if not ob:
            ob = core.OpeningBalance(base_date=base_date, bank_account_id=account.id, currency=account.currency, created_by=current_user.id); db.session.add(ob)
        ob.source_date = source_date; ob.calculation_method = method
        if final_data:
            tx = next(t for t in txs if str(t.id) == final_data["id"]); ob.source_balance = tx.balance; ob.opening_balance = tx.balance; ob.status = "confirmed"; ob.source_batch_id = tx.import_batch_id; ob.source_transaction_id = tx.id; confirmed += 1
        else:
            ob.source_balance = None; ob.opening_balance = None; ob.status = "needs_review"; ob.source_batch_id = None; ob.source_transaction_id = None; review += 1
    core.audit("opening_balances_generated", user=current_user, target_type="opening_balance", target_id=base_date.isoformat(), detail=f"confirmed={confirmed}, review={review}, manual_kept={manual_kept}")
    if core.commit_or_rollback(): flash(f"자동산출 완료: 자동확정 {confirmed} / 확인필요 {review} / Excel 확정값 보호 {manual_kept}", "success")
    else: flash("기초잔액 자동산출에 실패했습니다.", "error")
    return redirect(url_for("opening_balances"))


app.view_functions["opening_balances"] = login_required(opening_balances_view)
app.view_functions["generate_opening_balances"] = core.roles_required("admin", "editor")(generate_opening_balances_view)
app.add_url_rule("/data/opening-balances/template.xlsx", "opening_balance_template", login_required(opening_balance_template_view), methods=["GET"])
app.add_url_rule("/data/opening-balances/upload", "opening_balance_upload", core.roles_required("admin", "editor")(opening_balance_upload_view), methods=["POST"])
app.add_url_rule("/data/opening-balances/preview/<token>", "opening_balance_preview", login_required(opening_balance_preview_view), methods=["GET"])
app.add_url_rule("/data/opening-balances/confirm/<token>", "opening_balance_confirm", core.roles_required("admin", "editor")(opening_balance_confirm_view), methods=["POST"])
